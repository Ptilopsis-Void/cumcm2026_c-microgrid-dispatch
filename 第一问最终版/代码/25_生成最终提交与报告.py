import importlib
import openpyxl
import pandas as pd
from _comm import PROJECT_DIR as P, RESULT1_TEMPLATE, SUBMIT_RESULT1

SIX_DP = '0.000000'

def main():
    s = pd.read_csv(P/'模型结果/最终LP_原始最优解.csv')
    d = pd.read_csv(P/'处理后数据/附件一_第一问基础数据.csv')
    out = d.copy()
    mapping = {'grid_purchase_kw':'grid_purchase_power_kw', 'charge_power_kw':'charge_power_kw',
               'discharge_power_kw':'discharge_power_kw','pv_curtailment_kw':'unused_pv_power_kw',
               'energy_start_kwh':'energy_start_kwh','energy_end_kwh':'energy_end_kwh',
               'grid_purchase_energy_kwh':'grid_purchase_energy_kwh',
               'charge_input_energy_kwh':'charge_energy_kwh','discharge_output_energy_kwh':'discharge_energy_kwh',
               'period_purchase_cost_yuan':'period_purchase_cost_yuan'}
    for dst,src in mapping.items(): out[dst]=s[src]
    out['charge_state']=(s.charge_energy_kwh>1e-7).astype(int)
    out['battery_energy_increase_kwh']=s.charge_energy_kwh*.9
    out['battery_energy_decrease_kwh']=s.discharge_energy_kwh/.9
    out['power_balance_residual_kw']=s.power_balance_residual_kwh*6
    out['soc_transition_residual_kwh']=s.energy_transition_residual_kwh
    out.to_csv(P/'模型结果/第一问_最优调度结果.csv',index=False,encoding='utf-8-sig')
    baseline=importlib.import_module('08_计算无储能基准').compute_baseline(d)
    J=float(s.period_purchase_cost_yuan.sum());J0=baseline['J0'];Q=float(s.grid_purchase_energy_kwh.sum())
    costs=[('最低购电费用 J_1',J,'元'),('无储能购电费用 J_0',J0,'元'),('节省费用',J0-J,'元'),
           ('节省率',(J0-J)/J0*100,'%'),('总购电量',Q,'kWh'),('总充电输入电量',s.charge_energy_kwh.sum(),'kWh'),
           ('总放电输出电量',s.discharge_energy_kwh.sum(),'kWh'),('总未利用供能',s.unused_pv_energy_kwh.sum(),'kWh')]
    pd.DataFrame(costs,columns=['指标','数值','单位']).to_csv(P/'模型结果/第一问_费用汇总.csv',index=False,encoding='utf-8-sig')
    windows=[]
    for i in range(6):
        sub=s.iloc[i*24:(i+1)*24]
        windows.append([f'{i*4}:00-{(i+1)*4}:00',float(sub.charge_energy_kwh.sum()),float(sub.discharge_energy_kwh.sum())])
    def label(i):
        return f'{i//6}:{i%6*10:02d}-'+(f'{(i+1)//6}:{(i+1)%6*10:02d}' if i<143 else '0:00+1')
    wb=openpyxl.load_workbook(RESULT1_TEMPLATE)
    ws1=wb['计划购电量']
    ws2=wb['充放电量']
    for i in range(144):
        c=ws1.cell(i+2,1)
        c.value=label(i)
        c.number_format='@'
        v=ws1.cell(i+2,2)
        v.value=round(float(s.grid_purchase_energy_kwh.iloc[i]),6)
        v.number_format=SIX_DP
    for i,(name,ch,dis) in enumerate(windows):
        r=i+2
        ws2.cell(r,1).value=name
        for col,val in ((2,ch),(3,dis)):
            cell=ws2.cell(r,col)
            cell.value=round(float(val),6)
            cell.number_format=SIX_DP
    for r,stamp in ((2,'0:00'),(3,'24:00')):
        t=ws2.cell(r,4)
        t.value=stamp
        t.number_format='@'
        e=ws2.cell(r,5)
        e.value=6000
        e.number_format=SIX_DP
    wb.save(SUBMIT_RESULT1)
    lines=['# 第一问最终模型与计算结果','',
        '最终主模型为连续线性规划（LP）。全部提交数值、调度数据与图表统一来自 `模型结果/最终LP_原始最优解.csv`。历史 MILP 仅作最优值基准。','',
        '## 建模约定与模型','',
        '附件1的144个时间标签采用区间终点解释，区间内功率取该行值并视为恒定；这是本文的离散化约定。提交表时间标签调整为覆盖当天0:00—24:00，原始模板保持不变。',
        '设每段Δt=1/6小时，L、V为负荷与预测光伏电量，g、C、D为交流侧购电、充电、放电电量，E为内部储电量，U为未利用供能（供需不等式的非负松弛量）。',
        '- 目标：min Σ c_t g_t。',
        '- 供需平衡：g_t+V_t+D_t=L_t+C_t+U_t，g_t,U_t≥0。',
        '- 储能递推：E_t=E_(t−1)+0.9C_t−D_t/0.9。',
        '- 动作上限：0≤C_t,D_t≤5000/6 kWh。',
        '- 状态约束：1200≤E_t≤10800，E_0=E_144=6000 kWh。',
        '- 在非负电价和可丢弃剩余供能条件下，至少存在一组无同时充放电的最优解，本次解满足该性质，故不引入二进制变量。',
        '- 最优解U=0，也满足物理弃光范围U≤V。允许电网给储能充电，不向外网售电。','',
        '## 表1 指定时段购电量与全天合计','',
        '| 时间段 | 购电量（kWh） | 时间段 | 购电量（kWh） | 时间段 | 购电量（kWh） |','|---|---:|---|---:|---|---:|']
    for group in [[60,72,84],[96,108,120]]:
        cells=[]
        for i in group:cells.extend([label(i),f'{s.grid_purchase_energy_kwh.iloc[i]:.6f}'])
        lines.append('| '+' | '.join(cells)+' |')
    lines.extend([f'| 全天购电量 | {Q:.6f} | 全天购电费（元） | {J:.6f} | | |','',
        '## 表2 储能充放电量','',
        '| 时间段 | 充电量（kWh） | 放电量（kWh） | 时间段 | 充电量（kWh） | 放电量（kWh） |','|---|---:|---:|---|---:|---:|'])
    for i in range(0,6,2):
        cells=[]
        for name,ch,dis in windows[i:i+2]:cells.extend([name,f'{ch:.6f}',f'{dis:.6f}'])
        lines.append('| '+' | '.join(cells)+' |')
    lines.extend(['| 0:00储电量 | 6000.000000 | | 24:00储电量 | 6000.000000 | |','',
        '## 储能经济性','',f'无储能费用 {J0:.6f} 元；节省 {J0-J:.6f} 元，节省率 {(J0-J)/J0*100:.4f}%。','',
        '## 结果来源与解释','',
        '当前提交表与最终LP逐时一致，仅保留六位小数。存在等价最优轨迹不影响最优值，但正式文件统一使用上述唯一来源。',
        '对偶价格满足KKT条件；三种算法比较只说明本次结果的差异，不据此宣称对偶唯一。严格位于两阈值之间时静置，等于阈值时允许非零动作。',
        '价值函数使用后续时域LP重构并验证Bellman一致性，不声称采用独立DP求解器。最后时段真实内部折点为6000与6637.9122037 kWh，其他时刻仅报告采样斜率变化点数。',''])
    (P/'报告/第一问模型求解报告.md').write_text('\n'.join(lines))
    print(f'统一最终LP输出完成：J={J:.10f}, Q={Q:.10f}')

if __name__=='__main__':main()
