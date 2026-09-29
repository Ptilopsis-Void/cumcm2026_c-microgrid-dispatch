import hashlib
import importlib
from pathlib import Path
import tempfile
import numpy as np
import pandas as pd
from scipy.optimize import linprog
from scipy import sparse
from openpyxl import load_workbook
from _comm import PROJECT_DIR as P, ATTACHMENT1_PATH, RESULT1_TEMPLATE

def main():
    checks=[]
    def check(name,value,tol=1e-6):
        value=float(value);checks.append((name,value,tol,bool(np.isfinite(value) and value<=tol)))
    raw=pd.read_excel(ATTACHMENT1_PATH)
    d=pd.read_csv(P/'处理后数据/附件一_第一问基础数据.csv')
    s=pd.read_csv(P/'模型结果/最终LP_原始最优解.csv')
    for cn,en in [('电价','price_yuan_per_kwh'),('小区负载','load_kw'),('光伏发电预测功率','pv_forecast_kw')]:
        check(f'原始附件逐行一致_{cn}',np.max(abs(raw[cn]-d[en])),0)
    n=144;p=raw['电价'].to_numpy();L=raw['小区负载'].to_numpy()/6;V=raw['光伏发电预测功率'].to_numpy()/6
    c=np.r_[p/6,np.zeros(4*n)];A=np.zeros((2*n,5*n));b=np.r_[(L-V)*6,np.zeros(n)]
    for t in range(n):
        A[t,[t,n+t,2*n+t,3*n+t]]=[1,-1,1,-1]
        A[n+t,[n+t,2*n+t,4*n+t]]=[-.9/6,1/5.4,1]
        if t:A[n+t,4*n+t-1]=-1
        else:b[n+t]=6000
    bounds=[(0,None)]*n+[(0,5000)]*(2*n)+[(0,None)]*n+[(1200,10800)]*(n-1)+[(6000,6000)]
    r=linprog(c,A_eq=sparse.csr_matrix(A),b_eq=b,bounds=bounds,method='highs');assert r.success,r.message
    g,C,D,U=[s[k].to_numpy() for k in ['grid_purchase_energy_kwh','charge_energy_kwh','discharge_energy_kwh','unused_pv_energy_kwh']]
    es,ee=s.energy_start_kwh.to_numpy(),s.energy_end_kwh.to_numpy()
    check('独立功率LP与最终电量LP费用一致',abs(r.fun-p@g),1e-5)
    check('供需平衡',max(abs(g+V+D-L-C-U)))
    check('储能递推',max(abs(ee-es-.9*C+D/.9)))
    check('状态连续',max(abs(es[1:]-ee[:-1])))
    check('首末状态',max(abs(es[0]-6000),abs(ee[-1]-6000)))
    check('状态边界',max(0,1200-min(es.min(),ee.min()),max(es.max(),ee.max())-10800))
    check('充放电边界',max(0,-C.min(),-D.min(),C.max()-5000/6,D.max()-5000/6))
    check('购电和剩余非负',max(0,-g.min(),-U.min()))
    check('无同时充放电',int(((C>1e-7)&(D>1e-7)).sum()),0)
    m=importlib.import_module('13_建立并求解最终连续LP')
    cc,AA,bb,lb,ub=m.build_lp(p,L,V);xx=np.r_[g,C,D,U,es[0],ee]
    dual=pd.read_csv(P/'模型结果/最终LP_原始对偶数据.csv').raw_marginal.to_numpy()
    bd=pd.read_csv(P/'模型结果/最终LP_变量边界对偶数据.csv');lm=bd.lower_marginal.to_numpy();um=bd.upper_marginal.to_numpy();finite=np.isfinite(ub)
    check('完整KKT驻点',max(abs(cc-AA.T@dual-lm-um)),1e-8)
    check('下界互补',max(abs((xx-lb)*lm)),1e-6)
    check('上界互补',max(abs((ub[finite]-xx[finite])*um[finite])),1e-6)
    check('对偶符号',max(0,-lm.min(),um.max()),1e-8)
    check('强对偶间隙',abs(cc@xx-bb@dual-lb@lm-ub[finite]@um[finite]),1e-5)
    wb=load_workbook(P/'提交结果/result1.xlsx',data_only=True)
    wa=wb['计划购电量'];wb2=wb['充放电量'];wg=np.array([wa.cell(i+2,2).value for i in range(n)],float)
    check('提交表144段逐值与LP一致',max(abs(wg-g)),5.1e-7)
    check('提交费用舍入误差',abs(p@wg-p@g),1e-4)
    for i in range(6):
        check(f'四小时充电窗口{i+1}',abs(wb2.cell(i+2,2).value-C[i*24:i*24+24].sum()),5.1e-7)
        check(f'四小时放电窗口{i+1}',abs(wb2.cell(i+2,3).value-D[i*24:i*24+24].sum()),5.1e-7)
    check('提交表首末储能',max(abs(wb2['E2'].value-6000),abs(wb2['E3'].value-6000)),0)
    for i in range(n):
        expected=f'{i//6}:{i%6*10:02d}-'+(f'{(i+1)//6}:{(i+1)%6*10:02d}' if i<143 else '0:00+1')
        assert wa.cell(i+2,1).value==expected
    check('提交表时间标签全覆盖',0,0)
    plot=pd.read_csv(P/'模型结果/第一问_作图数据.csv')
    for col,vals in [('grid_purchase_kw',g*6),('charge_power_kw',C*6),('discharge_power_kw',D*6)]:
        check('作图数据与LP一致_'+col,max(abs(plot[col]-vals)),1e-6)
    ticks,labels=importlib.import_module('09_绘制第一问结果图')._time_ticks()
    check('图时间刻度',0 if list(ticks)==list(range(0,1441,120)) and labels==[f'{i:02d}:00' for i in range(0,25,2)] else 1,0)
    m21=importlib.import_module('21_构造连续状态价值函数');m22=importlib.import_module('22_校验连续DP与价值函数')
    knots=m21.build_cpl(p,L,V,144)
    check('末时段两个真实内部折点',0 if len(knots)==4 and abs(knots[1][0]-6000)<1e-8 and abs(knots[2][0]-(6000+(L[-1]-V[-1])/.9))<1e-8 else 1,0)
    with tempfile.TemporaryDirectory(prefix='q1_regression_') as tmp:
        tmp=Path(tmp);m22.OUT_MUTUAL=tmp/'mutual.csv';m22.OUT_BELLMAN=tmp/'bell.csv';m22.OUT_ACCEPT=tmp/'accept.csv'
        e=(knots[0][0]+knots[-1][0])/2
        traj=pd.DataFrame({'g':g,'C':C,'D':D,'U':U,'E_prev':es,'E_next':ee})
        ctx={'price':p,'L':L,'V':V,'n':L-V,'cpl_cache':{144:knots},'traj_df':traj,'E_end':ee[-1],'REP_TIMES':[(144,'23:50')]}
        healthy=m22.main(ctx,m21);check('验收正常输入通过',0 if healthy['all_pass'] else 1,0)
        orig=m22.bellman_rhs
        try:
            m22.bellman_rhs=lambda *args,**kwargs: float('inf')
            failed=m22.main(ctx,m21)
            check('无穷大残差必须拒绝',0 if not failed['all_pass'] else 1,0)
            m22.bellman_rhs=lambda *args,**kwargs: float('nan')
            failed=m22.main(ctx,m21)
            check('NaN残差必须拒绝',0 if not failed['all_pass'] else 1,0)
        finally:m22.bellman_rhs=orig
    for stem in ['最终LP_约束验收','第一问_对偶与KKT验收','连续DP_性质验收','同时充放电消去_验收结果']:
        path=P/f'模型结果/{stem}.csv'
        if not path.is_file():
            if stem == '同时充放电消去_验收结果':
                continue  # 开发期历史MILP基准未提供时，该可选互证不会生成
            raise FileNotFoundError(path)
        frame=pd.read_csv(path)
        if 'pass' in frame:check(stem,0 if frame['pass'].all() else 1,0)
        elif '结果' in frame:check(stem,0 if frame['结果'].isin(['通过','诊断（非失败项）']).all() else 1,0)
    result=pd.DataFrame(checks,columns=['验收项','误差或违例','容差','通过'])
    result.to_csv(P/'模型结果/最终版_独立验收.csv',index=False,encoding='utf-8-sig')
    result.to_csv(P/'模型结果/第一问_约束校验结果.csv',index=False,encoding='utf-8-sig')
    passed=int(result['通过'].sum());total=len(result)
    report=['# 第一问最终版验收记录','',f'独立验收：{passed}/{total} 项通过。','',
        '来源：原始附件1重新构建功率LP，完整KKT，最终LP与提交表及绘图数据逐项比对。',
        '回归测试：两小时刻度、末时段解析折点、正常验收以及无穷大/NaN残差拒绝。',
        f'最终购电费：{p@g:.10f} 元；总购电量：{g.sum():.10f} kWh。',
        '数值容差为本文自行设置，并非题目规定。视觉版式另行检查。','',
        '| 验收项 | 误差或违例 | 容差 | 通过 |','|---|---:|---:|---|']
    for name,val,tol,ok in checks:report.append(f'| {name} | {val:.3e} | {tol:.3e} | {"是" if ok else "否"} |')
    (P/'报告/第一问模型结果验收表.md').write_text('\n'.join(report)+'\n')
    print(f'独立验收 {passed}/{total} 项通过。')
    if passed!=total:raise RuntimeError(result[~result['通过']].to_string(index=False))

if __name__=='__main__':main()
