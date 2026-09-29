from pathlib import Path
import sys
from collections import Counter

import pandas as pd

sys.path.insert(0, str(Path(__file__).resolve().parent))
from _comm import (
    ATTACHMENT1_PATH, STRUCT_CSV, MAPPING_CSV, STAT_CSV, CHECK_MD,
    parse_time_label_to_minutes, label_to_text,
)

FIELD_MAP = [
    ("interval_index", "（派生）", "（无）", "时段编号 1..144", "模型正式输入"),
    ("original_time_label", "时间", "（无）", "10分钟区间终点标签（溯源用）", "模型正式输入"),
    ("interval_start", "（派生）", "（无）", "区间起点 HH:MM", "模型正式输入"),
    ("interval_end", "（派生）", "（无）", "区间终点 HH:MM", "模型正式输入"),
    ("delta_hours", "（常量）", "小时", "每时段时长 Δt = 1/6", "模型正式输入"),
    ("price_yuan_per_kwh", "电价", "元/kWh", "交易电价（模型参数 p_t）", "模型正式输入"),
    ("load_kw", "小区负载", "kW", "居民负荷功率（模型参数 L_t）", "模型正式输入"),
    ("pv_forecast_kw", "光伏发电预测功率", "kW", "光伏预测出力（模型参数 P_t^pv）", "模型正式输入"),
    ("net_load_kw", "（派生）", "kW", "load_kw - pv_forecast_kw", "描述性统计"),
    ("net_demand_kw", "（派生）", "kW", "max(net_load_kw, 0)", "描述性统计"),
    ("pv_surplus_kw", "（派生）", "kW", "max(-net_load_kw, 0)，静态差值，非弃光 W_t", "描述性统计"),
    ("load_energy_kwh", "（派生）", "kWh", "load_kw × Δt", "描述性统计"),
    ("pv_forecast_energy_kwh", "（派生）", "kWh", "pv_forecast_kw × Δt", "描述性统计"),
    ("net_demand_energy_kwh", "（派生）", "kWh", "net_demand_kw × Δt", "描述性统计"),
    ("pv_surplus_energy_kwh", "（派生）", "kWh", "pv_surplus_kw × Δt", "描述性统计"),
    ("g_t", "决策变量", "kWh", "外网计划购电量（由优化模型求解）", "后续模型输出"),
    ("C_t", "决策变量", "kWh", "交流侧充电量（由优化模型求解）", "后续模型输出"),
    ("D_t", "决策变量", "kWh", "交流侧放电量（由优化模型求解）", "后续模型输出"),
    ("E_t", "决策变量", "kWh", "储能电量（由优化模型求解）", "后续模型输出"),
    ("U_t", "决策变量", "kWh", "未利用供能（由优化模型求解，非 pv_surplus_kw）", "后续模型输出"),
]


def _fmt(x):
    if x is None:
        return ""
    return f"{x:.6g}"


def main():
    import openpyxl
    wb = openpyxl.load_workbook(ATTACHMENT1_PATH, data_only=True)
    sheet_names = wb.sheetnames
    ws = wb[sheet_names[0]]
    n_row = ws.max_row
    n_col = ws.max_column
    raw_headers = [ws.cell(1, c).value for c in range(1, n_col + 1)]

    time_cell_values = []
    time_cell_types = Counter()
    for r in range(2, n_row + 1):
        v = ws.cell(r, 1).value
        time_cell_values.append(v)
        time_cell_types[type(v).__name__] += 1

    df = pd.read_excel(ATTACHMENT1_PATH)
    cols = list(df.columns)
    n_data = len(df)

    struct_rows = [
        ["worksheet_name", sheet_names[0]],
        ["total_rows_with_header", n_row],
        ["data_rows", n_data],
        ["total_cols", n_col],
    ]
    for i, h in enumerate(raw_headers, 1):
        struct_rows.append([f"column_{i}", str(h)])
    struct_rows.append(["time_label_cell_type_dist", dict(time_cell_types)])
    pd.DataFrame(struct_rows, columns=["field", "value"]).to_csv(
        STRUCT_CSV, index=False, encoding="utf-8-sig")

    pd.DataFrame(FIELD_MAP, columns=["standard_field", "source", "unit", "description", "usage"]).to_csv(
        MAPPING_CSV, index=False, encoding="utf-8-sig")

    num_cols = [c for c in cols if pd.api.types.is_numeric_dtype(df[c])]
    stat_rows = []
    for c in num_cols:
        s = df[c]
        stat_rows.append([c, int(s.count()), _fmt(s.min()), _fmt(s.max()),
                          _fmt(s.mean()), _fmt(s.std(ddof=0))])
    pd.DataFrame(stat_rows, columns=["column", "count", "min", "max", "mean", "std"]).to_csv(
        STAT_CSV, index=False, encoding="utf-8-sig")

    head = df.head(5).to_string(index=False)
    tail = df.tail(5).to_string(index=False)

    miss = df.isna().sum()
    dup_rows = int(df.duplicated().sum())
    time_labels = list(df.iloc[:, 0])
    dup_label_set = sorted({str(t) for t in set(time_labels) if time_labels.count(t) > 1})

    minutes = [parse_time_label_to_minutes(v) for v in time_cell_values]
    is_monotonic = all(minutes[i] < minutes[i + 1] for i in range(len(minutes) - 1))
    diffs = [minutes[i + 1] - minutes[i] for i in range(len(minutes) - 1)]
    all_10min = all(d == 10 for d in diffs)
    is_144 = (n_data == 144)

    price = df["电价"]; load = df["小区负载"]; pv = df["光伏发电预测功率"]
    n_neg_price = int((price < 0).sum())
    n_neg_load = int((load < 0).sum())
    n_neg_pv = int((pv < 0).sum())

    pv_pos_idx = [i for i, v in enumerate(pv.values) if v > 0]
    pv_first = label_to_text(time_labels[pv_pos_idx[0]]) if pv_pos_idx else "无"
    pv_last = label_to_text(time_labels[pv_pos_idx[-1]]) if pv_pos_idx else "无"

    first_label = label_to_text(time_cell_values[0])
    last_label = label_to_text(time_cell_values[-1])

    lines = [
        "# 附件一原始数据检查报告",
        "",
        "## 1. 工作表名称", f"- {sheet_names}",
        "## 2. 表格行数和列数",
        f"- 含表头：{n_row} 行 × {n_col} 列", f"- 纯数据：{n_data} 行",
        "## 3. 原始列名",
    ]
    lines.append("| 序号 | 列名 |")
    lines.append("|---|---|")
    for i, h in enumerate(raw_headers, 1):
        lines.append(f"| {i} | {h} |")

    lines += [
        "", "## 4. 前5行与后5行",
        "### 前5行", "```", head, "```",
        "### 后5行", "```", tail, "```",
        "## 5. 数据类型",
    ]
    for c in cols:
        lines.append(f"- `{c}` : {df[c].dtype}")
    lines.append("")
    lines.append("> 时间列底层单元格类型混合，见下。")
    lines.append("")
    lines.append("### 时间列底层单元格类型统计")
    lines.append("| 类型 | 数量 |")
    lines.append("|---|---|")
    for t, c in sorted(time_cell_types.items()):
        lines.append(f"| {t} | {c} |")

    lines += [
        "", "## 6. 缺失值",
    ]
    for c in cols:
        lines.append(f"- `{c}` : {int(miss[c])}")

    lines += [
        "", "## 7. 重复行", f"- 完全重复行数：{dup_rows}",
        "## 8. 重复时间标签",
        f"- 重复标签：{dup_label_set if dup_label_set else '无'}",
        "## 9. 数值列统计（min / max / mean / std）",
        "| 列名 | 最小值 | 最大值 | 均值 | 标准差 |",
        "|---|---|---|---|---|",
    ]
    for c in num_cols:
        s = df[c]
        lines.append(f"| {c} | {_fmt(s.min())} | {_fmt(s.max())} | {_fmt(s.mean())} | {_fmt(s.std(ddof=0))} |")

    lines += [
        "", "## 10. 原始时间是否有序", f"- 严格单调递增：{is_monotonic}",
        "## 11. 相邻记录是否均为10分钟",
        f"- 相邻间隔全为10分钟：{all_10min}",
        f"- 首条标签：`{first_label}`，末条标签：`{last_label}`",
        "## 12. 是否恰好144个时段", f"- 恰好144行：{is_144}",
        "## 13. 负电价/负负荷/负光伏",
        f"- 负电价：{n_neg_price}", f"- 负负荷：{n_neg_load}", f"- 负光伏：{n_neg_pv}",
        "",
        "## 附：时间解释（证据）",
        f"- 首条标签 `{first_label}`（区间终点，非 00:00），末条 `{last_label}`（次日0点=当天24:00）。",
        f"- 光伏首个>0 标签：`{pv_first}`，末个>0 标签：`{pv_last}`。",
        "结论：标签为**区间终点**（详见“附件一时间标签映射说明.md”）。",
        "",
    ]
    CHECK_MD.write_text("\n".join(lines), encoding="utf-8")

    result = {
        "sheet_names": sheet_names,
        "data_rows": n_data,
        "cols": cols,
        "first_label": first_label,
        "last_label": last_label,
        "time_cell_types": dict(time_cell_types),
        "dup_rows": dup_rows,
        "dup_labels": len(dup_label_set),
        "n_missing": int(df.isna().sum().sum()),
        "n_neg_price": n_neg_price, "n_neg_load": n_neg_load, "n_neg_pv": n_neg_pv,
        "is_monotonic": is_monotonic, "all_10min": all_10min, "is_144": is_144,
        "pv_first": pv_first, "pv_last": pv_last,
    }
    print(f"[01] 检查完成：data_rows={n_data}, 工作表={sheet_names}, 列={cols}")
    print(f"[01] 已输出：{STRUCT_CSV.name}, {MAPPING_CSV.name}, {STAT_CSV.name}, {CHECK_MD.name}")
    return result


if __name__ == "__main__":
    main()
