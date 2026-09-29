from pathlib import Path
import importlib.util
import sys
import hashlib
import json
import numpy as np
import pandas as pd
import openpyxl

HERE = Path(__file__).resolve().parent

def load(name, alias):
    spec = importlib.util.spec_from_file_location(alias, HERE / name)
    module = importlib.util.module_from_spec(spec)
    sys.modules[alias] = module
    spec.loader.exec_module(module)
    return module

def main():
    DP = load('06_DP价值执行器.py', 'review_dp')
    P5 = load('05_求解日前计划.py', 'review_lp')
    P7 = load('07_执行器回测.py', 'review_exec')
    EX = load('10_生成并校验result2.py', 'review_export')
    C = DP.C
    path = C.RESULT_DIR / '第二问_执行器回测.npz'
    digest = hashlib.sha256(path.read_bytes()).hexdigest()
    base = np.load(C.BASE_MATRIX_NPZ)
    scen = np.load(C.SCENARIO_NPZ)
    bt = np.load(path)
    days = base['score_day_index']
    dates = [str(x) for x in base['dates']]
    price = base['price']
    for gaps in [np.array([[100.]]), np.array([[100., -200., 350., 0., 420.]])]:
        vf = DP.build_value_functions(gaps, np.zeros(gaps.shape[1]), np.ones(gaps.shape[1]))
        np.testing.assert_array_equal(vf['Rw'][0], vf['R'])
        assert vf['Rw'][0, -1] == C.E_MIN
    for starts in [[], [0, 5, 143], [0, 5, 10, 15, 20, 25, 30, 143]]:
        vals = np.zeros((1,144)); vals[0,starts] = np.arange(1,len(starts)+1)
        wb = openpyxl.Workbook(); ws = wb.active
        ws.cell(2,1).value = pd.Timestamp('2025-02-01').to_pydatetime()
        events = EX.merge_emergency(vals[0])
        for j in range(max(3,len(events))):
            if j < len(events):
                start,end,total = events[j]
                ws.cell(j+2,2).value = EX.seg_text(start,end)
                ws.cell(j+2,3).value = total
            else: ws.cell(j+2,2).value = ''
        assert EX.validate_emergency_sheet(ws,[0],['2025-02-01'],vals)[0]
        if events:
            saved = ws.cell(2,2).value; ws.cell(2,2).value = '00:00-24:00'
            assert not EX.validate_emergency_sheet(ws,[0],['2025-02-01'],vals)[0]
            ws.cell(2,2).value = saved; ws.cell(2,3).value += 1
            assert not EX.validate_emergency_sheet(ws,[0],['2025-02-01'],vals)[0]
    print('针对性回归检查通过：Rw 时间索引、零/3/8事件、午夜边界、错误边界及电量拒绝。', flush=True)
    rows = []
    lp = P5.DayPlanLP(price, nu=P7.NU, m=C.M_SCENARIOS)
    previous = C.E_INIT
    for k, day in enumerate(days):
        e0 = previous
        previous = float(bt['an_E'][day,-1])
        plan = bt['an_g'][day]
        solved = lp.solve(scen['scen_N'][day], e0)
        np.testing.assert_allclose(solved['g'],plan,atol=1e-6,rtol=0)
        co = (solved['b'] > P7.TOLL) & (solved['C'] > P7.TOLL)
        vf = DP.build_value_functions(scen['scen_N'][day],plan,price)
        lp_cost = 5 * (price @ solved['b'].T) - P7.NU * solved['E'][:,-1]
        evaluated = P7.eval_plan_causal(vf['Rw'],scen['scen_N'][day],plan,price,e0)
        rows.append((dates[day],int(co.any()),int(co.sum()),float(lp_cost.mean()),float(evaluated.mean()),float((evaluated-lp_cost).mean()),float(abs(evaluated-lp_cost).mean())))
        if (k+1)%50 == 0:print(f'诊断重算 {k+1}/{len(days)}',flush=True)
    dest = C.RESULT_DIR / '第二问_日前乐观偏差诊断.csv'
    headers = ['日期','是否出现同时购电与充电','同时发生情景时段数','日前情景目标均值_元','逐情景重评目标均值_元','有符号偏差_元','平均绝对差_元']
    C.write_csv_utf8_sig(dest,headers,rows)
    g,b,cc,d,e,u = [bt[key][days] for key in ['plan_g','dp_b','dp_C','dp_D','dp_E','dp_U']]
    prev = np.r_[C.E_INIT,e.ravel()[:-1]].reshape(e.shape)
    balance = float(abs(g+b+d-cc-u-base['net_load_energy_kwh'][days]).max())
    soc = float(abs(e-prev-C.ETA*cc+d/C.ETA).max())
    assert balance < 1e-7 and soc < 1e-7
    assert e.min() >= C.E_MIN-1e-7 and e.max() <= C.E_MAX+1e-7
    assert max(cc.max(),d.max()) <= C.S_PERIOD_KWH+1e-7 and np.minimum(cc,d).max() < 1e-7
    for date in ['2025-03-20','2025-06-21','2025-09-23','2025-12-21']:
        day = dates.index(date); k = list(days).index(day)
        vf = DP.build_value_functions(scen['scen_N'][day],bt['plan_g'][day],price)
        out = DP.dp_execute(vf['Hbar'],vf['R'],base['net_load_energy_kwh'][day],bt['plan_g'][day],prev[k,0])
        for key in ['C','D','b','E','U']: np.testing.assert_allclose(out[key],bt['dp_'+key][day],atol=1e-8,rtol=0)
    assert hashlib.sha256(path.read_bytes()).hexdigest() == digest
    summary = dict(days=len(days),diagnostic_signed_mean=float(np.mean([r[5] for r in rows])),diagnostic_abs_mean=float(np.mean([r[6] for r in rows])),co_days=sum(r[1] for r in rows),co_periods=sum(r[2] for r in rows),balance_residual=balance,soc_residual=soc,main_npz_sha256=digest,total_cost=float(((g+5*b)*price).sum()),regression_tests='通过',main_npz_unchanged=True)
    C.write_text_utf8(C.RESULT_DIR/'第二问_交付复核摘要.json',json.dumps(summary,ensure_ascii=False,indent=2))
    print(json.dumps(summary,ensure_ascii=False,indent=2),flush=True)
    return 0

if __name__ == '__main__':
    raise SystemExit(main())
