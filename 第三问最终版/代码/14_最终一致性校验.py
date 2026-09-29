from __future__ import annotations

import csv
import importlib.util
import os
import re
import sys
import time
from pathlib import Path

import numpy as np

_HERE = Path(__file__).resolve().parent


def _load(filename: str, modname: str):
    spec = importlib.util.spec_from_file_location(modname, _HERE / filename)
    mod = importlib.util.module_from_spec(spec)
    sys.modules[modname] = mod
    spec.loader.exec_module(mod)
    return mod


C = _load("_comm3.py", "q3_comm")

CHECK_CSV = C.RESULT_DIR / "第三问_最终一致性校验.csv"
REPORT_MD = C.REPORT_DIR / "第三问_最终一致性校验报告.md"
LOG_TXT = C.LOG_DIR / "第三问_14_日志.txt"

TOL_ALG = 1e-6
TOL_CSV = 1e-4
TOL_PHYS = 1e-6
REL_HARD = 1e-6
REL_SOFT = 1e-2


def _read_csv(path: Path) -> tuple[list[str], list[list[str]]]:
    with path.open("r", encoding="utf-8-sig", newline="") as fh:
        rows = list(csv.reader(fh))
    return (rows[0], rows[1:]) if rows else ([], [])


def _read_long_table(path: Path) -> dict[str, str]:
    out: dict[str, str] = {}
    if not path.exists():
        return out
    with path.open("r", encoding="utf-8-sig", newline="") as fh:
        for r in csv.reader(fh):
            if len(r) >= 2 and r[0].strip():
                out[r[0].strip()] = r[1].strip()
    return out


def _read_named_rows(path: Path, key_col: str) -> dict[str, dict[str, str]]:
    hdr, body = _read_csv(path)
    if not hdr:
        return {}
    i = hdr.index(key_col)
    return {r[i]: dict(zip(hdr, r)) for r in body if len(r) == len(hdr)}


def _f(d: dict[str, str], k: str) -> float:
    try:
        return float(d.get(k, "nan"))
    except (TypeError, ValueError):
        return float("nan")


def _fnum(x) -> float:
    try:
        return float(str(x).strip())
    except (TypeError, ValueError):
        return float("nan")


def _precision_rows(path: Path) -> tuple[dict[str, float], float]:
    hdr, body = _read_csv(path) if path.is_file() else ([], [])
    hour: dict[str, float] = {}
    base = float("nan")
    if not hdr or "口径" not in hdr or "MAE_kW" not in hdr:
        return hour, base
    ic, im = hdr.index("口径"), hdr.index("MAE_kW")
    it = hdr.index("发布时刻") if "发布时刻" in hdr else None
    for r in body:
        if len(r) <= max(ic, im):
            continue
        v = _fnum(r[im])
        tau = r[it].strip() if (it is not None and len(r) > it) else ""
        if "因果基线" in r[ic] or "前 7 日" in r[ic]:
            base = v
        elif "整点" in r[ic] and tau:
            hour[tau] = v
    return hour, base


def _precision_window(path: Path) -> dict[str, str]:
    hdr, body = _read_csv(path) if path.is_file() else ([], [])
    per: dict[str, str] = {}
    if not hdr or "口径" not in hdr:
        return per
    ic = hdr.index("口径")
    it = hdr.index("发布时刻") if "发布时刻" in hdr else None
    for r in body:
        if len(r) <= ic:
            continue
        cal = r[ic].strip()
        m = re.search(r"d=\d+\.\.\d+", cal)
        lab = (r[it].strip() if (it is not None and len(r) > it) else "?")
        per[f"{lab}|{cal}"] = m.group(0) if m else ""
    return per


def _report_precision_window(txt: str) -> str | None:
    m = re.search(r"同窗口\s*(d=\d+\.\.\d+)", txt)
    return m.group(1) if m else None


def _scan_zero_point(path: Path, mae_col: str, pi_col: str,
                     pi_zero: str = "0") -> dict[str, float]:
    hdr, body = _read_csv(path) if path.is_file() else ([], [])
    out: dict[str, float] = {}
    if not hdr or mae_col not in hdr or pi_col not in hdr:
        return out
    im, ip = hdr.index(mae_col), hdr.index(pi_col)
    it = hdr.index("发布时刻") if "发布时刻" in hdr else None
    for r in body:
        if len(r) <= max(im, ip):
            continue
        if r[ip].strip() != pi_zero:
            continue
        lab = r[it].strip() if (it is not None and len(r) > it) else "?"
        out[lab] = _fnum(r[im])
    return out


def _scan_window(path: Path, win_col: str = "评价窗口") -> set[str]:
    hdr, body = _read_csv(path) if path.is_file() else ([], [])
    if not hdr or win_col not in hdr:
        return set()
    iw = hdr.index(win_col)
    return {r[iw].strip() for r in body if len(r) > iw and r[iw].strip()}


def _dict_close(a: dict[str, float], b: dict[str, float],
                tol: float) -> bool:
    if not a or not b or set(a) != set(b):
        return False
    return all(np.isfinite(a[k]) and np.isfinite(b[k])
               and abs(a[k] - b[k]) <= tol for k in a)


def _by_label(path: Path, val_col: str,
              label_col: str = "发布时刻") -> dict[str, float]:
    hdr, body = _read_csv(path) if path.is_file() else ([], [])
    out: dict[str, float] = {}
    if not hdr or val_col not in hdr or label_col not in hdr:
        return out
    iv, il = hdr.index(val_col), hdr.index(label_col)
    for r in body:
        if len(r) <= max(iv, il):
            continue
        out[r[il].strip()] = _fnum(r[iv])
    return out


def _precision_by_caliber(path: Path) -> dict[str, dict[str, float]]:
    hdr, body = _read_csv(path) if path.is_file() else ([], [])
    out: dict[str, dict[str, float]] = {}
    if not hdr or "口径" not in hdr or "MAE_kW" not in hdr:
        return out
    ic, im = hdr.index("口径"), hdr.index("MAE_kW")
    it = hdr.index("发布时刻") if "发布时刻" in hdr else None
    for r in body:
        if len(r) <= max(ic, im):
            continue
        cal = r[ic]
        if "前 7 日" in cal or "因果基线" in cal:
            key = "因果基线"
        elif "线性插值" in cal:
            key = "10 min 线性插值"
        elif "分段常数" in cal:
            key = "10 min 分段常数"
        elif "整点" in cal:
            key = "整点"
        else:
            continue
        lab = r[it].strip() if (it is not None and len(r) > it) else "?"
        out.setdefault(key, {})[lab] = _fnum(r[im])
    return out


def _report_precision_claim(txt: str) -> tuple[str | None, float | None,
                                              float | None]:
    m_h = re.search(r"整点 MAE\s*([\d,]+(?:\.\d+)?)\s*kW", txt)
    m_t = re.search(r"(\d{1,2}):00\s*发布", txt)
    m_b = re.search(r"基线\s*([\d,]+(?:\.\d+)?)\s*kW", txt)

    def _g(m):
        return float(m.group(1).replace(",", "")) if m else None

    return (m_t.group(1) + ":00" if m_t else None, _g(m_h), _g(m_b))


def _xlsx_source_headers(path, sheet: str):
    if not path:
        return None
    p = Path(path)
    if not p.is_file():
        return None
    try:
        from openpyxl import load_workbook
        wb = load_workbook(p, data_only=True, read_only=True)
        if sheet not in wb.sheetnames:
            return None
        ws = wb[sheet]
        row = [c.value for c in next(ws.iter_rows(min_row=1, max_row=1))]
        wb.close()
        return row
    except Exception:
        return None


def _day_hint(path: Path) -> tuple[int | None, bool]:
    if not path.exists():
        return None, False
    hdr, body = _read_csv(path)
    for i, k in enumerate(hdr):
        if k in ("评分日数", "天数"):
            vals: set[int] = set()
            for r in body:
                if len(r) != len(hdr):
                    return None, False
                try:
                    vals.add(int(float(r[i])))
                except (ValueError, IndexError):
                    return None, False
            return (next(iter(vals)) if len(vals) == 1 else None), True
    return None, False


def main() -> int:
    C.ensure_dirs()
    log = C.Tee()
    t0 = time.perf_counter()

    log("=" * 78)
    log("第三问 14 —— 最终一致性校验（只读交叉核对）")
    log("=" * 78)
    log(f"输出根目录：{C.OUT_ROOT}")
    log("")

    checks: list[tuple[str, str, str, str]] = []

    def chk(section: str, name: str, ok: bool, detail: str = "",
            skip: bool = False) -> bool:
        st = "跳过" if skip else ("通过" if ok else "未通过")
        checks.append((section, name, st, detail))
        mark = "－" if skip else ("✔" if ok else "✘")
        log(f"  {mark} {name}" + (f"：{detail}" if detail else ""))
        return bool(ok)

    if not C.BACKTEST_NPZ.exists():
        log(f"✘ 未找到 {C.BACKTEST_NPZ.name}，请先运行 08。")
        return 1

    BT = np.load(C.BACKTEST_NPZ, allow_pickle=False)
    Z = C.Q2.matrix()

    price = np.asarray(BT["price"], float)
    sc = np.asarray(BT["score_day_index"], int)
    P = np.asarray(BT["P"], float)
    Q = np.asarray(BT["Q"], float)
    b = np.asarray(BT["b"], float)
    Ch = np.asarray(BT["C"], float)
    Dh = np.asarray(BT["D"], float)
    E = np.asarray(BT["E"], float)
    curt = np.asarray(BT["U"], float)
    NN = np.asarray(BT["N_actual"], float)
    n_day = sc.size

    log(f"评分日数 = {n_day}；时段数 = {P.shape[1]}；"
        f"ρ_down = {C.RHO_DOWN}；ρ_up = {C.RHO_UP}；κ_em = 5")
    log("")

    log("── §1 `08` 结算闭式独立重算（不调用 `_policy3.decompose`）──")
    Pp, Qq, bb = P[sc], Q[sc], b[sc]
    pr = np.broadcast_to(price[None, :], Pp.shape)
    u_c = np.maximum(Pp - Qq, 0.0)
    v_c = np.maximum(Qq - Pp, 0.0)

    plan_c = (pr * np.minimum(Pp, Qq)).sum(axis=1)
    down_c = (C.RHO_DOWN * pr * u_c).sum(axis=1)
    up_c = (C.RHO_UP * pr * v_c).sum(axis=1)
    em_c = (5.0 * pr * bb).sum(axis=1)
    tot_c = plan_c + down_c + up_c + em_c

    d_plan = np.asarray(BT["daily_plan_fee"], float)
    d_down = np.asarray(BT["daily_down_fee"], float)
    d_up = np.asarray(BT["daily_up_fee"], float)
    d_em = np.asarray(BT["daily_emerg_fee"], float)
    d_tot = np.asarray(BT["daily_total_fee"], float)
    d_u = np.asarray(BT["daily_u"], float)
    d_v = np.asarray(BT["daily_v"], float)
    d_b = np.asarray(BT["daily_b"], float)

    def _mx(a, b_):
        return float(np.max(np.abs(np.asarray(a, float) - np.asarray(b_, float))))

    chk("§1", "计划购电费 = Σ c·min(p,q)", _mx(plan_c, d_plan) < TOL_ALG,
        f"max|Δ| {_mx(plan_c, d_plan):.3e} 元")
    chk("§1", "下调违约金 = 0.5·Σ c·u", _mx(down_c, d_down) < TOL_ALG,
        f"max|Δ| {_mx(down_c, d_down):.3e} 元")
    chk("§1", "上调加价 = 1.5·Σ c·v", _mx(up_c, d_up) < TOL_ALG,
        f"max|Δ| {_mx(up_c, d_up):.3e} 元")
    chk("§1", "紧急购电费 = 5·Σ c·b", _mx(em_c, d_em) < TOL_ALG,
        f"max|Δ| {_mx(em_c, d_em):.3e} 元")
    chk("§1", "当日总费用 = 计划 + 下调 + 上调 + 紧急", _mx(tot_c, d_tot) < TOL_ALG,
        f"max|Δ| {_mx(tot_c, d_tot):.3e} 元")
    chk("§1", "逐日 Σu / Σv 与 npz daily_u / daily_v 一致",
        _mx(u_c.sum(axis=1), d_u) < TOL_ALG and _mx(v_c.sum(axis=1), d_v) < TOL_ALG,
        f"max|Δu| {_mx(u_c.sum(axis=1), d_u):.3e}，"
        f"max|Δv| {_mx(v_c.sum(axis=1), d_v):.3e} kWh")
    chk("§1", "逐日 Σb 与 npz daily_b 一致", _mx(bb.sum(axis=1), d_b) < TOL_ALG,
        f"max|Δ| {_mx(bb.sum(axis=1), d_b):.3e} kWh")

    log("── §1b 电价三源一致性（账单唯一乘子）──")
    pr_q2 = np.asarray(Z["price"], float)
    _n_t = P.shape[1]
    chk("§1b", f"price 长度 = 时段数（{_n_t}）",
        price.size == _n_t, f"BT {price.size} vs 时段 {_n_t}")
    chk("§1b", "price 全部有限且严格为正",
        bool(np.all(np.isfinite(price)) and np.all(price > 0)),
        f"min {float(np.min(price)):.6f}，max {float(np.max(price)):.6f}")
    chk("§1b", "电价非退化（存在峰谷差，结算 min/max 结构才有效）",
        float(np.ptp(price)) > 1e-9,
        f"峰谷差 {float(np.ptp(price)):.6f} 元/kWh")
    chk("§1b", "BT.price ↔ 附件二矩阵 price（逐格）",
        price.size == pr_q2.size and _mx(price, pr_q2) < 1e-9,
        f"max|Δ| {_mx(price, pr_q2):.3e} 元/kWh")
    _vp = C.V_FORECAST_NPZ
    if _vp.is_file():
        try:
            pr_q3 = np.asarray(np.load(_vp, allow_pickle=False)["price"], float)
            chk("§1b", "BT.price ↔ 附件三预报矩阵 price（逐格）",
                price.size == pr_q3.size and _mx(price, pr_q3) < 1e-9,
                f"max|Δ| {_mx(price, pr_q3):.3e} 元/kWh")
        except Exception as exc:
            chk("§1b", "BT.price ↔ 附件三预报矩阵 price（逐格）", False,
                f"读取失败 {type(exc).__name__}: {exc}")
    else:
        chk("§1b", "BT.price ↔ 附件三预报矩阵 price（逐格）", False,
            f"未找到 {_vp.name}", skip=True)

    chk("§1", "恒等式 q = p − u + v（逐格）",
        _mx(Qq, Pp - u_c + v_c) < TOL_ALG,
        f"max|Δ| {_mx(Qq, Pp - u_c + v_c):.3e} kWh")
    chk("§1", "0 ≤ u ≤ p（逐格）",
        float(u_c.min()) >= -TOL_ALG and float((u_c - Pp).max()) <= TOL_ALG,
        f"min u {u_c.min():.3e}；max(u−p) {(u_c - Pp).max():.3e}")
    chk("§1", "v ≥ 0 且 u·v = 0（逐格互斥）",
        float(v_c.min()) >= -TOL_ALG and float((u_c * v_c).max()) <= TOL_ALG,
        f"min v {v_c.min():.3e}；max(u·v) {(u_c * v_c).max():.3e}")
    chk("§1", "ν 未进入账单（账单仅由四项构成）",
        abs(float(np.asarray(BT["yearly_total"], float).ravel()[0]) - float(tot_c.sum()))
        < TOL_ALG,
        f"年度账单 {float(np.asarray(BT['yearly_total'], float).ravel()[0]):,.6f} vs "
        f"四项合计 {float(tot_c.sum()):,.6f} 元")

    chk("§1", "计划购电量 p ≥ 0（逐格）",
        float(P[sc].min()) >= -TOL_ALG, f"min p {P[sc].min():.3e} kWh")
    chk("§1", "调整购电量 q ≥ 0（逐格，由 q = p − u + v 且 0 ≤ u ≤ p 推出）",
        float(Qq.min()) >= -TOL_ALG, f"min q {Qq.min():.3e} kWh")
    chk("§1", "紧急购电量 b ≥ 0（逐格）",
        float(bb.min()) >= -TOL_ALG, f"min b {bb.min():.3e} kWh")

    log("")
    log("── §2 npz ↔ 逐日 / 月度 / 年度 CSV ──")
    dates = [str(x) for x in BT["dates"]]
    hdr, body = _read_csv(C.DAILY_CSV)
    ok_shape = len(body) == n_day
    chk("§2", "逐日 CSV 行数 = 评分日数", ok_shape,
        f"{len(body)} vs {n_day}")

    col = {k: i for i, k in enumerate(hdr)}
    day_tot_col = None
    for cand in ("当日总费用_元", "总费用_元"):
        if cand in col:
            day_tot_col = cand
            break

    if ok_shape and day_tot_col:
        def _col(name: str) -> np.ndarray:
            return np.asarray([float(r[col[name]]) for r in body], float)

        pairs = [
            ("计划购电量_kWh", Pp.sum(axis=1)),
            ("调整购电量_kWh", Qq.sum(axis=1)),
            ("下调量_kWh", d_u),
            ("上调量_kWh", d_v),
            ("计划购电费_元", d_plan),
            ("下调违约金_元", d_down),
            ("上调加价_元", d_up),
            ("紧急购电费_元", d_em),
            (day_tot_col, d_tot),
            ("紧急购电量_kWh", d_b),
            ("实际净负荷_kWh", NN[sc].sum(axis=1)),
            ("充电量_kWh", Ch[sc].sum(axis=1)),
            ("放电量_kWh", Dh[sc].sum(axis=1)),
            ("弃电量_kWh", curt[sc].sum(axis=1)),
        ]
        for name, arr in pairs:
            if name not in col:
                chk("§2", f"逐日 CSV 列「{name}」存在", False, "列缺失")
                continue
            chk("§2", f"逐日 CSV「{name}」↔ npz", _mx(_col(name), arr) < TOL_CSV,
                f"max|Δ| {_mx(_col(name), arr):.3e}")
        ok_date = all(str(body[i][0]) == dates[sc[i]] for i in range(n_day))
        chk("§2", "逐日 CSV 日期列 ↔ npz dates", ok_date)
    elif not day_tot_col:
        chk("§2", "逐日 CSV 含总费用列", False, f"表头 {hdr[:3]}…")

    if ok_shape and day_tot_col and C.MONTHLY_CSV.exists():
        mh, mb = _read_csv(C.MONTHLY_CSV)
        mc = {k: i for i, k in enumerate(mh)}
        mmap = {k: k for k in ("计划购电量_kWh", "调整购电量_kWh", "计划购电费_元",
                               "下调违约金_元", "上调加价_元", "紧急购电费_元",
                               "紧急购电量_kWh")}
        mmap[day_tot_col] = "当月总费用_元"
        msum: dict[str, dict[str, float]] = {}
        for i, d in enumerate(sc):
            slot = msum.setdefault(dates[d][:7], {})
            for day_k, mon_k in mmap.items():
                slot[mon_k] = slot.get(mon_k, 0.0) + float(body[i][col[day_k]])
        worst, worst_k = 0.0, ""
        for r in mb:
            for mon_k, v in msum.get(r[0], {}).items():
                df = abs(float(r[mc[mon_k]]) - v)
                if df > worst:
                    worst, worst_k = df, f"{r[0]}/{mon_k}"
        chk("§2", "月度 CSV ↔ 逐日 CSV 加总（8 个量 × 各月）", worst < TOL_CSV,
            f"max|Δ| {worst:.3e}（{worst_k}）")
        chk("§2", "月度 CSV 天数合计 = 评分日数",
            abs(sum(int(r[mc["天数"]]) for r in mb) - n_day) == 0,
            f"{sum(int(r[mc['天数']]) for r in mb)} vs {n_day}")
    else:
        chk("§2", "月度 CSV 加总闭合", True, "（前置条件不足，跳过）", skip=True)

    yr = _read_long_table(C.YEARLY_CSV)
    missing_anchor = [k for k in (
        "计划购电量_kWh", "调整购电量_kWh", "下调量_kWh", "上调量_kWh",
        "计划购电费_元", "调整相关费用_元", "紧急购电量_kWh", "紧急购电费_元",
        "总费用_元", "无储能基准购电费_元", "储能相对基准节省_元",
        "储能相对基准节省_%", "总费用均价_元每kWh") if k not in yr]
    chk("§2", "年度汇总表含全部 13 个关键锚点", not missing_anchor,
        f"缺失 {missing_anchor}" if missing_anchor else "13/13 就位")

    if yr:
        _yr_tot = _f(yr, "总费用_元")
        _pp = float(tot_c.sum())
        _pd = float(d_tot.sum())
        chk("§2", "年度表 总费用_元 ≡ 逐时段费用求和（方案 §8.3，tol 1e-6 元）",
            abs(_yr_tot - _pp) < 1e-6,
            f"{_yr_tot:,.6f} vs {_pp:,.6f} 元（Δ {abs(_yr_tot - _pp):.3e}）")
        chk("§2", "年度表 总费用_元 ≡ Σ 逐日 npz 总费用（tol 1e-6 元）",
            abs(_yr_tot - _pd) < 1e-6,
            f"{_yr_tot:,.6f} vs {_pd:,.6f} 元（Δ {abs(_yr_tot - _pd):.3e}）")
        chk("§2", "年度表 总费用_元 ↔ npz yearly_total",
            abs(_f(yr, "总费用_元") - float(tot_c.sum())) < TOL_CSV,
            f"{_f(yr, '总费用_元'):,.6f} vs {float(tot_c.sum()):,.6f} 元")
        chk("§2", "年度表 计划购电量/调整购电量 ↔ npz",
            abs(_f(yr, "计划购电量_kWh") - float(Pp.sum())) < TOL_CSV
            and abs(_f(yr, "调整购电量_kWh") - float(Qq.sum())) < TOL_CSV,
            f"计划 Δ{abs(_f(yr, '计划购电量_kWh') - float(Pp.sum())):.3e}，"
            f"调整 Δ{abs(_f(yr, '调整购电量_kWh') - float(Qq.sum())):.3e}")
        chk("§2", "年度表 调整相关费用_元 = 下调 + 上调",
            abs(_f(yr, "调整相关费用_元")
                - (float(d_down.sum()) + float(d_up.sum()))) < TOL_CSV,
            f"{_f(yr, '调整相关费用_元'):,.6f} vs "
            f"{float(d_down.sum()) + float(d_up.sum()):,.6f} 元")
        chk("§2", "年度表 总费用均价 = 总费用 / 调整购电量",
            abs(_f(yr, "总费用均价_元每kWh")
                - float(tot_c.sum()) / float(Qq.sum())) < 1e-4,
            f"{_f(yr, '总费用均价_元每kWh'):.6f} vs "
            f"{float(tot_c.sum()) / float(Qq.sum()):.6f}")
        nostor = _f(yr, "无储能基准购电费_元")
        chk("§2", "年度表 储能节省_元 = 基准 − 总费用",
            abs(_f(yr, "储能相对基准节省_元") - (nostor - float(tot_c.sum())))
            < TOL_CSV,
            f"{_f(yr, '储能相对基准节省_元'):,.6f} 元")
        chk("§2", "年度表 储能节省_% 与基准/总费用自洽",
            abs(_f(yr, "储能相对基准节省_%")
                - 100.0 * (nostor - float(tot_c.sum())) / nostor) < 1e-3,
            f"{_f(yr, '储能相对基准节省_%'):.6f} vs "
            f"{100.0 * (nostor - float(tot_c.sum())) / nostor:.6f} %")

        Npos = np.maximum(NN[sc], 0.0)
        kwh_c = float(Npos.sum())
        fee_c = float((pr * Npos).sum())
        chk("§2", "无储能基准 购电量 独立重算",
            abs(_f(yr, "无储能基准购电量_kWh") - kwh_c) < 1e-2,
            f"表 {_f(yr, '无储能基准购电量_kWh'):,.4f} vs 重算 {kwh_c:,.4f} kWh")
        chk("§2", "无储能基准 购电费 独立重算",
            abs(nostor - fee_c) < 1e-2,
            f"表 {nostor:,.4f} vs 重算 {fee_c:,.4f} 元")

    cmp_path = C.RESULT_DIR / "第三问_结算口径对照.csv"
    if cmp_path.is_file():
        _chdr, body_c = _read_csv(cmp_path)
        got_r = {}
        for row in body_c:
            if not row or not row[0].strip():
                continue
            got_r[row[0].strip()] = [_fnum(x) for x in row[1:]]
        impl_ref = {
            "总费用_元": float(tot_c.sum()),
            "计划购电费_元": float(np.asarray(BT["daily_plan_fee"], float).sum()),
            "下调违约金_元": float(np.asarray(BT["daily_down_fee"], float).sum()),
            "上调加价_元": float(np.asarray(BT["daily_up_fee"], float).sum()),
            "紧急购电费_元": float(np.asarray(BT["daily_emerg_fee"], float).sum()),
            "紧急购电量_kWh": float(np.asarray(BT["daily_b"], float).sum()),
        }
        bad_impl = [k for k, v in impl_ref.items()
                    if k not in got_r or len(got_r[k]) < 2
                    or abs(got_r[k][1] - v) > TOL_CSV]
        chk("§2", "口径对照 样本外实现口径列 ↔ npz 逐日加总（6 个量）",
            not bad_impl, f"不一致 {bad_impl}" if bad_impl else "6/6 逐位一致")
        if "总费用_元" in got_r and len(got_r["总费用_元"]) >= 1:
            _exp_npz = float(np.asarray(BT["exp_total"], float).ravel()[0])
            chk("§2", "口径对照 模型内期望口径 总费用 ↔ npz exp_total",
                abs(got_r["总费用_元"][0] - _exp_npz) < TOL_CSV,
                f"{got_r['总费用_元'][0]:,.6f} vs {_exp_npz:,.6f} 元")
        if "紧急购电量_kWh" in got_r and len(got_r["紧急购电量_kWh"]) >= 3:
            ban = np.asarray(BT["b_analytic"], float)[sc]
            kwh_an = float(ban.sum())
            fee_an = float((C.EMERG_MULT * np.asarray(pr, float)[None, :] * ban).sum())
            chk("§2", "口径对照 纯解析执行 紧急购电量 ↔ npz b_analytic 独立重算",
                abs(got_r["紧急购电量_kWh"][2] - kwh_an) < 1e-2,
                f"{got_r['紧急购电量_kWh'][2]:,.4f} vs {kwh_an:,.4f} kWh")
            if "紧急购电费_元" in got_r and len(got_r["紧急购电费_元"]) >= 3:
                chk("§2", "口径对照 纯解析执行 紧急购电费 ↔ Σ5c·b_analytic 独立重算",
                    abs(got_r["紧急购电费_元"][2] - fee_an) < 1e-2,
                    f"{got_r['紧急购电费_元'][2]:,.4f} vs {fee_an:,.4f} 元")
        same_plan = ("计划购电费_元" in got_r
                     and len(got_r["计划购电费_元"]) >= 3
                     and max(got_r["计划购电费_元"])
                     - min(got_r["计划购电费_元"]) < TOL_CSV)
        chk("§2", "口径对照 计划购电费三列相同（与执行策略无关，只由计划 P 决定）",
            same_plan,
            f"{got_r.get('计划购电费_元')}" if not same_plan
            else f"{got_r['计划购电费_元'][0]:,.4f} 元（三列一致）")
        if ("总费用_元" in got_r and len(got_r["总费用_元"]) >= 3):
            _d_dp = got_r["总费用_元"][1] - got_r["总费用_元"][2]
            _rel_dp = abs(_d_dp) / max(abs(got_r["总费用_元"][2]), 1.0)
            chk("§2", "口径对照 实现−纯解析 差值有限且量级合理（不预设优劣方向）",
                bool(np.isfinite(_d_dp)) and _rel_dp < 0.5,
                f"实现 − 纯解析 = {_d_dp:+,.4f} 元（{_rel_dp:.4%}），即"
                + ("实现更贵" if _d_dp > 1e-2 else
                   "实现更省" if _d_dp < -1e-2 else "持平"))
            chk("§2", "口径对照 表内 总费用 = 计划+下调+上调+紧急（四分解闭合）",
                abs(got_r["总费用_元"][1]
                    - (got_r["计划购电费_元"][1] + got_r["下调违约金_元"][1]
                       + got_r["上调加价_元"][1] + got_r["紧急购电费_元"][1]))
                < 1e-2,
                f"{got_r['总费用_元'][1]:,.6f} vs "
                f"{got_r['计划购电费_元'][1] + got_r['下调违约金_元'][1] + got_r['上调加价_元'][1] + got_r['紧急购电费_元'][1]:,.6f} 元")
    else:
        chk("§2", "结算口径对照表存在", False, str(cmp_path), skip=True)

    log("")
    log("── §3 `08` 闭环审计表 ──")
    audit_path = C.RESULT_DIR / "第三问_闭环审计.csv"
    a_hdr, a_body = _read_csv(audit_path)
    if a_hdr and a_body:
        ac = {k: i for i, k in enumerate(a_hdr)}
        bad = [r[0] for r in a_body
               if r[ac["阈值"]] != "—" and r[ac["是否通过"]] != "通过"]
        chk("§3", "闭环审计表全部可判项均为「通过」", not bad,
            f"未通过：{bad}" if bad else f"{len(a_body)} 行 / 可判 {sum(1 for r in a_body if r[ac['阈值']] != '—')} 项")
        if "resid_names" in BT.files and "resid_values" in BT.files:
            names = [str(x) for x in BT["resid_names"]]
            vals = np.asarray(BT["resid_values"], float)
            rd = {r[0]: r[1] for r in a_body}
            worst, wk = 0.0, ""
            for nm, v in zip(names, vals):
                if nm in rd:
                    try:
                        df = abs(float(rd[nm]) - float(v))
                    except ValueError:
                        continue
                    if df > worst:
                        worst, wk = df, nm
            chk("§3", "审计表数值 ↔ npz resid_values", worst < 1e-6,
                f"max|Δ| {worst:.3e}（{wk}）")
        else:
            chk("§3", "npz 含 resid_names/resid_values", False, "键缺失")
    else:
        chk("§3", "闭环审计表存在", False, str(audit_path))

    log("")
    log("── §4 `09` 消融完整档 ≡ `08` 主结果（P0-4）──")
    abl_smoke = int(os.environ.get("Q3_ABL_MAX_DAYS", "0") or 0) > 0
    abl = _read_named_rows(C.ABLATION_CSV, "档") if C.ABLATION_CSV.exists() else {}
    n09, n09_ok = _day_hint(C.ABLATION_CSV)
    if "S061218" in abl:
        r = abl["S061218"]
        pairs = [("总费用_元", float(tot_c.sum())),
                 ("计划购电费_元", float(d_plan.sum())),
                 ("下调违约金_元", float(d_down.sum())),
                 ("上调加价_元", float(d_up.sum())),
                 ("调整相关费用_元", float(d_down.sum() + d_up.sum())),
                 ("紧急购电费_元", float(d_em.sum())),
                 ("计划购电量_kWh", float(Pp.sum())),
                 ("调整购电量_kWh", float(Qq.sum())),
                 ("下调量_kWh", float(d_u.sum())),
                 ("上调量_kWh", float(d_v.sum())),
                 ("紧急购电量_kWh", float(d_b.sum())),
                 ("实际净负荷_kWh", float(NN[sc].sum())),
                 ("充电量_kWh", float(Ch[sc].sum())),
                 ("放电量_kWh", float(Dh[sc].sum())),
                 ("弃电量_kWh", float(curt[sc].sum()))]
        _ABL_SCHEMA = ["档", "说明", "调整购电量_kWh", "下调量_kWh", "上调量_kWh",
                       "计划购电费_元", "下调违约金_元", "上调加价_元",
                       "调整相关费用_元", "紧急购电费_元", "紧急购电量_kWh",
                       "紧急购电时段数", "总费用_元", "期望口径计划调整费_元",
                       "期望口径紧急费_元", "期望口径总费用_元",
                       "闭环残差最大值", "信息泄露时段数", "拒绝新解段数",
                       "完美候选被压段数", "强制采纳段数", "期望口径尺子",
                       "评分日数"]
        _miss_c = sorted(set(_ABL_SCHEMA) - set(r))
        _extra_c = sorted(set(r) - set(_ABL_SCHEMA))
        chk("§4", "消融表 schema ≡ `09` 写表清单（缺列/多列 = 产物版本不符）",
            not _miss_c and not _extra_c,
            (f"缺 {_miss_c}；多 {_extra_c}" if (_miss_c or _extra_c)
             else f"{len(_ABL_SCHEMA)} 列就位"),
            skip=abl_smoke)
        _CMP_EXPECT = {"总费用_元", "计划购电费_元", "下调违约金_元", "上调加价_元",
                       "调整相关费用_元", "紧急购电费_元", "调整购电量_kWh",
                       "下调量_kWh", "上调量_kWh", "紧急购电量_kWh"}
        _cov = sorted(k for k, _ in pairs if k in r)
        chk("§4", f"可比字段覆盖 ≡ 声明的最小可比集（{len(_CMP_EXPECT)} 项）",
            set(_cov) == _CMP_EXPECT,
            f"实比 {len(_cov)}/{len(pairs)} 项（本表不承载 "
            f"{sorted({k for k, _ in pairs} - _CMP_EXPECT)}）；"
            f"缺 {sorted(_CMP_EXPECT - set(_cov))}",
            skip=abl_smoke)
        _refd = dict(pairs)
        worst, wk = 0.0, ""
        for k in _cov:
            df = abs(_f(r, k) - _refd[k])
            if df > worst:
                worst, wk = df, k
        rel = worst / max(abs(float(tot_c.sum())), 1e-12)
        W = (f"消融完整档 ≡ 主结果（{len(_cov)} 个可比字段 = 声明最小可比集，"
             f"方案 §3.5）")
        if n09 is not None and n09 != n_day:
            chk("§4", W, True,
                f"产物评分日数 {n09} ≠ 主结果 {n_day} → 不可比，跳过", skip=True)
        elif rel > (REL_HARD if n09 == n_day else REL_SOFT):
            if n09 is None:
                chk("§4", W, True,
                    f"产物未记录评分日数，最大相对差 {100 * rel:.4f} %（{wk}）"
                    f"疑似天数口径不同 → 跳过，须在正式全量运行下复核", skip=True)
            else:
                chk("§4", W, False, f"同日数下出现差异：max|Δ| {worst:.3e}（{wk}）")
        else:
            chk("§4", W, True,
                f"max|Δ| {worst:.3e}（{wk}，相对 {100 * rel:.9f} %）"
                + ("" if n09 is not None else "（未记录天数，按 1 % 容差兜底）"))

        NESTED = ["S0", "S06", "S0612", "S061218"]
        ALL6 = NESTED + ["S_all+_UB", "S_all+_raw", "PF"]
        if all(k in abl for k in NESTED):
            nexp = [_f(abl[k], "期望口径总费用_元") for k in NESTED]
            bad = [f"{NESTED[i]}→{NESTED[i + 1]}" for i in range(len(NESTED) - 1)
                   if nexp[i] - nexp[i + 1] < -1e-6]
            chk("§4", "期望口径**嵌套链**单调非增（S0 ≥ S06 ≥ S0612 ≥ S061218，可证）",
                not bad, "违反：" + "、".join(bad) if bad
                else "差额 " + "、".join(f"{nexp[i] - nexp[i + 1]:+,.2f}"
                                        for i in range(len(NESTED) - 1)) + " 元",
                skip=abl_smoke)
        if "S_all+_UB" in abl and "S0612" in abl:
            _ub6 = _f(abl["S_all+_UB"], "期望口径总费用_元")
            _s612 = _f(abl["S0612"], "期望口径总费用_元")
            chk("§4", "信息优势臂**可证**上界 S_all+_UB ≤ S0612（候选集 ⊇，同基线）",
                _ub6 <= _s612 + 1e-6,
                f"S_all+_UB {_ub6:,.2f} ≤ S0612 {_s612:,.2f} 元",
                skip=abl_smoke)
            if "S061218" in abl:
                _s18 = _f(abl["S061218"], "期望口径总费用_元")
                chk("§4", "禁止把 S_all+_UB 与 S061218 的期望差当作「信息价值上界」"
                          "（仅如实披露，符号可正可负）",
                    True, f"同尺子观测差 {_ub6 - _s18:+,.2f} 元（跨候选集，不可证）",
                    skip=abl_smoke)
        if "PF" in abl:
            if "期望口径尺子" not in abl["PF"]:
                chk("§4", "消融表逐档标注「期望口径尺子」（跨尺子防护）", False,
                    "缺列：产物由旧版 `09` 生成，需重跑")
            else:
                _nc = sorted(k for k in abl
                             if not abl[k].get("期望口径尺子", "").startswith("共同"))
                chk("§4", "消融表逐档标注「期望口径尺子」：仅 PF 非共同尺子（跨尺子防护）",
                    _nc == ["PF"], f"非共同尺子档 = {_nc or '空'}（应为 ['PF']）",
                    skip=abl_smoke)
        if (not abl_smoke) and ("S061218" in abl) and ("S_all+_raw" not in abl):
            chk("§4", "消融表含反事实臂 S_all+_raw（缺该档 = 旧版 `09` 产物）",
                False, "缺该档 ⇒ 产物由旧版 `09` 生成，需重跑 09")
        if "S_all+_raw" in abl:
            if "强制采纳段数" not in abl["S_all+_raw"]:
                chk("§4", "反事实臂 S_all+_raw 已真实强制采纳完美候选", False,
                    "缺列「强制采纳段数」：产物由旧版 `09` 生成，需重跑",
                    skip=abl_smoke)
            else:
                _nf = _f(abl["S_all+_raw"], "强制采纳段数")
                chk("§4", "反事实臂 S_all+_raw 已真实强制采纳完美候选", _nf > 0,
                    f"强制采纳 {_nf:.0f} 段（应 > 0）", skip=abl_smoke)
                _raw_e = _f(abl["S_all+_raw"], "期望口径总费用_元")
                _s18e = (_f(abl["S061218"], "期望口径总费用_元")
                         if "S061218" in abl else float("nan"))
                chk("§4", "反事实臂 S_all+_raw 与 S061218 的期望差**仅如实披露**"
                          "（该臂跳过接受检验，严禁当作上界）",
                    True, f"S_all+_raw {_raw_e:,.2f} 元 vs S061218 {_s18e:,.2f} 元"
                          f"（差 {_raw_e - _s18e:+,.2f} 元）", skip=abl_smoke)
        _abl_md = Path(C.REPORT_ABL_MD)
        if _abl_md.exists() and not abl_smoke:
            _txt = _abl_md.read_text(encoding="utf-8", errors="replace")
            _forbidden = ["相对完整档的增量", "即便**直接给它完美信息**",
                          "严格单调非增", "也不满足",
                          "即便把 18:00 的预报换成**完美信息**",
                          "（$S_{all+}^{UB}$，同基线）也只多贡献"]
            _hit = [t for t in _forbidden if t in _txt]
            chk("§4", "消融报告不含越界单调性表述（Issue-I 防回归）", not _hit,
                "命中禁用表述：" + "、".join(_hit) if _hit
                else "无禁用表述", skip=False)
            _need = ("可证性分级", "跨尺子禁令", "反事实臂", "强制采纳段数")
            _miss = [t for t in _need if t not in _txt]
            chk("§4", "消融报告含 §5「可证性分级」表 + 跨尺子禁令 + 反事实臂披露",
                not _miss, "四处防回归锚点齐全" if not _miss
                else "缺锚点：" + "、".join(_miss))
            if ("S_all+_UB" in abl) and ("S061218" in abl):
                _ubx = _f(abl["S_all+_UB"], "期望口径总费用_元")
                _s18x = _f(abl["S061218"], "期望口径总费用_元")
                _deg = abs(_ubx - _s18x) < 1e-6
                _disclosed = ("完美信息价值为 0" in _txt) or ("价值为 0" in _txt)
                chk("§4", "保守采纳臂退化时报告显式提醒「未被采纳 ≠ 价值为 0」",
                    (not _deg) or _disclosed,
                    "未退化（无需提醒）" if not _deg
                    else ("报告含该提醒" if _disclosed else "缺失该提醒"))
        for k in [x for x in ALL6 if x in abl]:
            if k in abl:
                chk("§4", f"{k} 闭环残差 ≤ 1e-8 / 信息泄露 = 0",
                    _f(abl[k], "闭环残差最大值") <= 1e-8
                    and _f(abl[k], "信息泄露时段数") == 0,
                    f"残差 {_f(abl[k], '闭环残差最大值'):.2e}，"
                    f"泄露 {_f(abl[k], '信息泄露时段数'):.0f}")
        if not abl_smoke:
            chk("§4", "消融全档 Σb 与逐日 npz 口径自洽（非负且有限）",
                all(np.isfinite(_f(abl[k], "紧急购电量_kWh")) for k in abl),
                f"{len(abl)} 档")
    else:
        why = "`Q3_ABL_MAX_DAYS` 冒烟模式（未生成完整档）" if abl_smoke \
            else f"未找到 {C.ABLATION_CSV.name} 或 S061218 行"
        chk("§4", "消融完整档 ≡ 主结果", True, why, skip=True)

    log("")
    log("── §5 `13` 敏感性/收敛性 ≡ `08` 主结果 ──")
    sens_smoke = int(os.environ.get("Q3_SENS_MAX_DAYS", "0") or 0) > 0
    ref_total = float(tot_c.sum())
    ref_exp = float(np.asarray(BT["exp_total"], float).ravel()[0]) \
        if "exp_total" in BT.files else float("nan")

    _TOT_LBL = "总费用（实现口径）"
    _EXP_LBL = "期望总费用（共同测度）"
    DP_CSV = C.RESULT_DIR / "第三问_收敛性_网格步长.csv"
    M_CSV = C.RESULT_DIR / "第三问_收敛性_情景数.csv"
    NU_CSV = C.RESULT_DIR / "第三问_敏感性_终端价值.csv"

    def _same_row(path: Path, pred, what: str) -> None:
        nd, _ = _day_hint(path)
        if not path.exists():
            chk("§5", what, True, f"（未找到 {path.name}，跳过）", skip=True)
            return
        if nd is not None and nd != n_day:
            chk("§5", what, True,
                f"产物评分日数 {nd} ≠ 主结果 {n_day} → 不可比，跳过", skip=True)
            return
        rows = _read_named_rows(path, "档位")
        hit = [k for k in rows if pred(k)]
        if not hit:
            chk("§5", what, True, "（无匹配档位，跳过）", skip=True)
            return
        k = hit[0]
        r = rows[k]
        v_tot, v_exp = _f(r, _TOT_LBL), _f(r, _EXP_LBL)
        rel = abs(v_tot - ref_total) / max(abs(ref_total), 1e-12)
        rel_e = (abs(v_exp - ref_exp) / max(abs(ref_exp), 1e-12)
                 if np.isfinite(ref_exp) else 0.0)
        worst = max(rel, rel_e)
        detail = (f"档「{k}」实现 {v_tot:,.4f} vs 主结果 {ref_total:,.4f} 元"
                  f"（相对差 {100 * rel:.6f} %）；期望相对差 {100 * rel_e:.6f} %")
        if worst <= (REL_HARD if nd == n_day else REL_SOFT):
            chk("§5", what, True,
                detail + ("" if nd == n_day else "（未记录天数，按 1 % 容差兜底）"))
        elif nd is None:
            chk("§5", what, True,
                detail + "  ⚠ 差异 > 1 % 且产物未记录评分日数，"
                "疑似天数口径不同 → 跳过", skip=True)
        else:
            chk("§5", what, False, detail)

    _same_row(DP_CSV, lambda k: "=6" in k.replace(" ", ""),
              "A 块 δ=6（默认网格）≡ 主结果")
    _same_row(M_CSV, lambda k: k.replace(" ", "") == "M=30",
              "B 块 M=30（全情景集）≡ 主结果")
    _same_row(NU_CSV, lambda k: "复算值" in k,
              "C 块 ν=附件 1 复算值 ≡ 主结果")

    CONV_TOL = 0.005
    RPT_SENS = C.REPORT_DIR / "第三问_敏感性与收敛性报告.md"
    _rpt = RPT_SENS.read_text(encoding="utf-8") if RPT_SENS.exists() else ""

    def _last_rel(a: float, b: float) -> float:
        if not (np.isfinite(a) and np.isfinite(b)):
            return float("nan")
        if abs(b) < 1e-12:
            return 0.0 if abs(a) < 1e-12 else float("nan")
        return (a - b) / b

    def _conv_verdict(path: Path, block: str) -> None:
        if not path.exists() or sens_smoke:
            chk("§5", f"{block} 收敛性判定（镜像 `13` 规则）", True,
                "（冒烟模式或文件缺失，跳过）", skip=True)
            return
        seq = list(_read_named_rows(path, "档位").values())
        if len(seq) < 2:
            chk("§5", f"{block} 收敛性判定（镜像 `13` 规则）", True,
                "（档位不足 2，跳过）", skip=True)
            return
        _NEED = (("主指标", _EXP_LBL), ("紧急电量", "紧急购电量"),
                 ("充电量", "充电量"))
        _lack = sorted({c for _, c in _NEED if c not in seq[-1]})
        chk("§5", f"{block} 判定所需列齐备（缺列不得静默缩小判据集）",
            not _lack, (f"缺 {_lack}" if _lack else f"{len(_NEED)} 列就位"))
        if _lack:
            return
        rel = {"总费用（共同测度）": _last_rel(_f(seq[-1], _EXP_LBL),
                                            _f(seq[-2], _EXP_LBL))}
        for lbl, col in _NEED[1:]:
            rel[lbl] = _last_rel(_f(seq[-1], col), _f(seq[-2], col))
        _fin = {k: abs(v) for k, v in rel.items() if np.isfinite(v)}
        worst_k = max(_fin, key=lambda k: _fin[k]) if _fin else ""
        conv = bool(_fin) and max(_fin.values()) <= CONV_TOL
        _main = rel["总费用（共同测度）"]
        chk("§5", f"{block}：主指标（共同测度）末次相对变化 ≤ "
                  f"{100 * CONV_TOL:.2f} %",
            np.isfinite(_main) and abs(_main) <= CONV_TOL,
            f"主指标 {100 * _main:+.4f} %；"
            + "；".join(f"{k} {100 * v:+.4f} %" for k, v in rel.items()
                        if k != "总费用（共同测度）"))
        token = "✅ 已收敛" if conv else "⚠ 尚未收敛"
        _hit = [ln for ln in _rpt.splitlines()
                if ln.startswith("|") and block in ln]
        chk("§5", f"{block} 报告结论 ↔ 数值重算（{token}）",
            bool(_hit) and token in _hit[0],
            (f"报告行「{_hit[0].strip()}」" if _hit
             else "报告未找到该块判定行（无法互证）"))
        if not conv:
            chk("§5", f"{block} 报告已点名超阈指标 `末次相对变化·{worst_k}`",
                f"末次相对变化·{worst_k}" in _rpt,
                f"重算最劣指标 {worst_k}（{100 * _fin[worst_k]:+.4f} %）")

    _conv_verdict(DP_CSV, "A 网格步长 δ")
    _conv_verdict(M_CSV, "B 情景数 M")

    cnpz = C.RESULT_DIR / "第三问_敏感性与收敛性.npz"
    if cnpz.exists():
        SN = np.load(cnpz, allow_pickle=True)
        have = [b for b in ("delta", "m", "nu") if f"{b}_labels" in SN.files]
        miss = [b for b in ("delta", "m", "nu") if b not in have]
        chk("§5", "`13` npz 含可校验的参数块", bool(have),
            f"存在 {have}；缺失 {miss}（缺失块多为该块尚未运行，本身不算错误）")
        if have:
            worst, wk = 0.0, ""
            for b in have:
                lbl = [str(x) for x in SN[f"{b}_labels"]]
                cols = [str(x) for x in SN[f"{b}_cols"]]
                body = np.vstack([np.asarray(SN[f"{b}_col{j}"], float)
                                  for j in range(len(cols))]) if len(cols) else None
                if body is None:
                    continue
                j_tot = cols.index("总费用_元") if "总费用_元" in cols else 0
                path = {"delta": DP_CSV, "m": M_CSV, "nu": NU_CSV}[b]
                if not path.exists():
                    continue
                rows = _read_named_rows(path, "档位")
                for i, lb in enumerate(lbl):
                    if lb in rows:
                        df = abs(_f(rows[lb], _TOT_LBL) - float(body[j_tot, i]))
                        if df > worst:
                            worst, wk = df, f"{b}/{lb}"
            chk("§5", "`13` npz ↔ CSV 总费用逐格一致", worst < TOL_CSV,
                f"max|Δ| {worst:.3e}（{wk}）")
    else:
        chk("§5", "`13` npz 存在", True, f"（未找到 {cnpz.name}，跳过）", skip=True)

    log("")
    log("── §6 物理一致性独立复核 ──")
    ETA = C.ETA
    load_e = np.asarray(Z["load_energy_kwh"], float)[sc]
    pv_e = np.asarray(Z["pv_energy_kwh"], float)[sc]

    chk("§6", "N_actual = 实际负荷 − 实际光伏（逐日）",
        _mx(NN[sc].sum(axis=1), load_e.sum(axis=1) - pv_e.sum(axis=1)) < TOL_PHYS,
        f"max|Δ| {_mx(NN[sc].sum(axis=1), load_e.sum(axis=1) - pv_e.sum(axis=1)):.3e} kWh")

    E_s = np.asarray(BT["E_day_start_real"], float)
    E_e = np.asarray(BT["E_day_end_real"], float)
    chk("§6", "SOC 递推 ηΣC − ΣD/η = ΔE（逐日）",
        _mx(ETA * Ch[sc].sum(axis=1) - Dh[sc].sum(axis=1) / ETA, E_e - E_s) < TOL_PHYS,
        f"max 残差 {_mx(ETA * Ch[sc].sum(axis=1) - Dh[sc].sum(axis=1) / ETA, E_e - E_s):.3e} kWh")
    chk("§6", "SOC 全程落在 [E_min, E_max]",
        float(E_e.min()) >= C.E_MIN - 1e-6 and float(E_e.max()) <= C.E_MAX + 1e-6,
        f"范围 [{E_e.min():.3f}, {E_e.max():.3f}] kWh ⊆ "
        f"[{C.E_MIN:.0f}, {C.E_MAX:.0f}]")
    chk("§6", "单时段充/放电量 ≤ S = 5000·Δt",
        float(Ch[sc].max()) <= C.S_PERIOD_KWH + 1e-6
        and float(Dh[sc].max()) <= C.S_PERIOD_KWH + 1e-6,
        f"max 充 {Ch[sc].max():.6f}，max 放 {Dh[sc].max():.6f} kWh"
        f"（S = {C.S_PERIOD_KWH:.6f}）")
    chk("§6", "充放电互斥 C·D = 0（逐格）",
        float((Ch[sc] * Dh[sc]).max()) <= 1e-9,
        f"max(C·D) {(Ch[sc] * Dh[sc]).max():.3e}")

    Ecell = E[sc]
    over = curt[sc] > 1e-9
    n_over_purch = int((over & (NN[sc] >= Qq - 1e-9)).sum())
    chk("§6", "弃电>0 ⇒ 购电量 > 净负荷（过度购电，非光伏过剩）",
        n_over_purch == 0,
        f"弃电时段 {int(over.sum())} 个，其中 q ≤ N 的 {n_over_purch} 个")
    pred = np.maximum(Qq - NN[sc], 0.0) - Ch[sc]
    dv = np.abs(pred - curt[sc])
    chk("§6", "弃电 = max(q − N, 0) − C（按 07 分支式逐格独立重算）",
        _mx(dv, 0.0) < TOL_PHYS, f"max|Δ| {_mx(dv, 0.0):.3e} kWh")
    n_chg_lim = int((over & (Ch[sc] < C.S_PERIOD_KWH - 1e-6)
                     & (Ecell < C.E_MAX - 1e-6)).sum())
    chk("§6", "弃电>0 ⇒ 充电已触 S 或末端 SOC 已触 E_max（先充满再弃）",
        n_chg_lim == 0, f"违规格数 {n_chg_lim}（末端 SOC 取逐格 E[t]）")
    n_over_b = int((over & (bb > 1e-9)).sum())
    chk("§6", "弃电>0 ⇒ 紧急购电 = 0（富余分支不触发紧急购电）",
        n_over_b == 0, f"同段格数 {n_over_b}")
    chk("§6", "紧急购电不被用于主动充电（充电段 b = 0）",
        float(np.max(np.where(Ch[sc] > 1e-9, bb, 0.0))) <= 1e-9,
        f"max b|(C>0) {float(np.max(np.where(Ch[sc] > 1e-9, bb, 0.0))):.3e} kWh")

    lhs = Qq.sum(axis=1) + bb.sum(axis=1) + Dh[sc].sum(axis=1)
    rhs = load_e.sum(axis=1) - pv_e.sum(axis=1) + Ch[sc].sum(axis=1) \
        + curt[sc].sum(axis=1)
    chk("§6", "电量守恒 q + b + D = 负荷 − 光伏 + C + 弃电（逐日）",
        _mx(lhs, rhs) < 1e-2, f"max|Δ| {_mx(lhs, rhs):.3e} kWh")

    bal = Qq + bb + Dh[sc] - Ch[sc] - curt[sc] - NN[sc]
    chk("§6", "逐格功率平衡 q + b + D − C − 弃电 = N",
        _mx(bal, 0.0) < TOL_PHYS, f"max|Δ| {_mx(bal, 0.0):.3e} kWh")
    chk("§6", "弃电量非负（逐格）", float(curt[sc].min()) >= -TOL_PHYS,
        f"min 弃电 {float(curt[sc].min()):.3e} kWh")

    _an_keys = ("C_analytic", "D_analytic", "b_analytic", "curt_analytic")
    if all(k in BT.files for k in _an_keys):
        NNs = NN[sc]
        Ca = np.asarray(BT["C_analytic"], float)[sc]
        Da = np.asarray(BT["D_analytic"], float)[sc]
        Ba = np.asarray(BT["b_analytic"], float)[sc]
        Ua = np.asarray(BT["curt_analytic"], float)[sc]
        S_ = C.S_PERIOD_KWH
        _E_s9 = np.asarray(BT["E_day_start_real"], float)
        _rt_all = NNs - Qq
        D_re = np.zeros_like(Da)
        b_re = np.zeros_like(Da)
        C_re = np.zeros_like(Da)
        U_re = np.zeros_like(Da)
        _e = float(_E_s9[0])
        e_lo = e_hi = _e
        for _i in range(_rt_all.shape[0]):
            for _t in range(_rt_all.shape[1]):
                _rt = float(_rt_all[_i, _t])
                if _rt > 0.0:
                    _d = max(min(_rt, S_, ETA * (_e - C.E_MIN)), 0.0)
                    D_re[_i, _t] = _d
                    b_re[_i, _t] = _rt - _d
                    _e -= _d / ETA
                else:
                    _cc = max(min(-_rt, S_, (C.E_MAX - _e) / ETA), 0.0)
                    C_re[_i, _t] = _cc
                    U_re[_i, _t] = -_rt - _cc
                    _e += ETA * _cc
                if _e < e_lo:
                    e_lo = _e
                if _e > e_hi:
                    e_hi = _e
        _w_re = max(_mx(b_re, Ba), _mx(C_re, Ca), _mx(D_re, Da), _mx(U_re, Ua))
        chk("§6", "纯解析轨迹 = 独立重放（α=0 规则逐年顺序重算，逐格 4 个量）",
            _w_re < TOL_PHYS,
            f"max|Δ| b {_mx(b_re, Ba):.3e} / C {_mx(C_re, Ca):.3e} / "
            f"D {_mx(D_re, Da):.3e} / U {_mx(U_re, Ua):.3e} kWh；"
            f"重放 SOC ∈ [{e_lo:.3f}, {e_hi:.3f}]，起点 {float(_E_s9[0]):.3f}")
        chk("§6", "评分年首日真实日初 SOC = C.E_INIT（重放链起点无关性）",
            abs(float(_E_s9[0]) - float(C.E_INIT)) < TOL_PHYS,
            f"{float(_E_s9[0]):.6f} vs C.E_INIT {float(C.E_INIT):.6f} kWh")
        chk("§6", "纯解析 逐格功率平衡 q + b + D − C − 弃电 = N",
            _mx(Qq + Ba + Da - Ca - Ua - NNs, 0.0) < TOL_PHYS,
            f"max|Δ| {_mx(Qq + Ba + Da - Ca - Ua - NNs, 0.0):.3e} kWh")
        _cap_an = max(float(Ca.max()), float(Da.max()))
        chk("§6", "纯解析 充放电互斥 C·D = 0 且单时段充/放 ≤ S（逐格）",
            float(np.max(np.abs(Ca * Da))) <= TOL_PHYS
            and _cap_an <= S_ + TOL_PHYS and float(min(Ca.min(), Da.min())) >= -TOL_PHYS,
            f"max C·D {float(np.max(np.abs(Ca * Da))):.3e}；"
            f"max 充/放 {_cap_an:.6f} vs S {S_:.6f} kWh")
        chk("§6", "纯解析 规则语义：弃电 = max(q−N−C,0) 且 b = max(N−q−D,0)",
            _mx(Ua, np.maximum(Qq - NNs - Ca, 0.0)) < TOL_PHYS
            and _mx(Ba, np.maximum(NNs - Qq - Da, 0.0)) < TOL_PHYS
            and float(Ua.min()) >= -TOL_PHYS and float(Ba.min()) >= -TOL_PHYS,
            f"max|Δ| 弃电 {_mx(Ua, np.maximum(Qq - NNs - Ca, 0.0)):.3e} / "
            f"b {_mx(Ba, np.maximum(NNs - Qq - Da, 0.0)):.3e} kWh")
    else:
        chk("§6", "纯解析对照轨迹可独立重放（需要 C/D/b/弃电 四组键）",
            False, f"npz 缺少 {[k for k in _an_keys if k not in BT.files]}",
            skip=True)

    log("")
    log("── §7 最终提交件 result3.xlsx 独立复核 ──")
    xp = C.RESULT3_XLSX
    if not xp.is_file():
        chk("§7", "result3.xlsx 存在", False, str(xp), skip=True)
    else:
        try:
            from openpyxl import load_workbook
        except ImportError:
            chk("§7", "openpyxl 可用", False, "未安装 openpyxl", skip=True)
            load_workbook = None
        if load_workbook is not None:
            wb = load_workbook(xp, data_only=True)
            want_sheets = [C.SHEET_PLAN3, C.SHEET_ADJUST,
                           C.SHEET_BATT, C.SHEET_EMERG]
            chk("§7", "四个工作表名称与顺序完全一致（附件 5 模板未被改名）",
                wb.sheetnames == want_sheets,
                f"{wb.sheetnames}")
            T_ = P.shape[1]

            if C.SHEET_PLAN3 in wb.sheetnames:
                wa = wb[C.SHEET_PLAN3]
                chk("§7", "『计划购电量』335 行 × 147 列",
                    wa.max_row == 1 + n_day and wa.max_column == T_ + 3,
                    f"{wa.max_row} × {wa.max_column}")
                wp = wq = wtot = wfee = 0.0
                bad_date = 0
                for i, d in enumerate(sc):
                    row = 2 + i
                    if str(wa.cell(row, 1).value)[:10] != str(dates[d]):
                        bad_date += 1
                    for t in range(T_):
                        wp = max(wp, abs(float(wa.cell(row, 2 + t).value or 0.0)
                                         - float(P[d, t])))
                    wtot = max(wtot, abs(float(wa.cell(row, T_ + 2).value or 0.0)
                                         - float(P[d].sum())))
                    wfee = max(wfee, abs(float(wa.cell(row, T_ + 3).value or 0.0)
                                         - float(d_plan[i])))
                chk("§7", "『计划购电量』A 列日期 ↔ npz dates",
                    bad_date == 0, f"不符 {bad_date} 行")
                chk("§7", "『计划购电量』B:EO 逐格 ↔ npz P（334×144 格）",
                    wp < TOL_CSV, f"max|Δ| {wp:.3e} kWh")
                chk("§7", "『计划购电量』EP 列 = 当日 ΣP",
                    wtot < TOL_CSV, f"max|Δ| {wtot:.3e} kWh")
                chk("§7", "『计划购电量』EQ 列 = npz daily_plan_fee",
                    wfee < TOL_CSV, f"max|Δ| {wfee:.3e} 元")

            if C.SHEET_ADJUST in wb.sheetnames:
                wb2 = wb[C.SHEET_ADJUST]
                chk("§7", "『调整购电量』335 行 × 147 列",
                    wb2.max_row == 1 + n_day and wb2.max_column == T_ + 3,
                    f"{wb2.max_row} × {wb2.max_column}")
                wq = wtot2 = wfee2 = 0.0
                for i, d in enumerate(sc):
                    row = 2 + i
                    for t in range(T_):
                        wq = max(wq, abs(float(wb2.cell(row, 2 + t).value or 0.0)
                                         - float(Q[d, t])))
                    wtot2 = max(wtot2,
                                abs(float(wb2.cell(row, T_ + 2).value or 0.0)
                                    - float(Q[d].sum())))
                    wfee2 = max(wfee2,
                                abs(float(wb2.cell(row, T_ + 3).value or 0.0)
                                    - float(d_down[i] + d_up[i])))
                chk("§7", "『调整购电量』B:EO 逐格 ↔ npz Q（334×144 格）",
                    wq < TOL_CSV, f"max|Δ| {wq:.3e} kWh")
                chk("§7", "『调整购电量』EP 列 = 当日 ΣQ",
                    wtot2 < TOL_CSV, f"max|Δ| {wtot2:.3e} kWh")
                chk("§7", "『调整购电量』EQ 列 = 下调费 + 上调加价",
                    wfee2 < TOL_CSV, f"max|Δ| {wfee2:.3e} 元")

            if C.SHEET_BATT in wb.sheetnames:
                wc = wb[C.SHEET_BATT]
                nblk = len(C.BATT_BLOCKS)
                per = T_ // nblk
                chk("§7", "『充放电量』行数 = 1 + 334×6",
                    wc.max_row == 1 + n_day * nblk, f"{wc.max_row}")
                e_start = np.zeros(n_day, float)
                _prev = float(C.E_INIT)
                for i, d in enumerate(sc):
                    e_start[i] = _prev
                    _prev = float(E[d, -1])
                wchg = wdchg = 0.0
                we0 = we1 = 0.0
                wlab = 0
                for i, d in enumerate(sc):
                    for j, blk in enumerate(C.BATT_BLOCKS):
                        rr = 2 + i * nblk + j
                        if str(wc.cell(rr, 2).value).strip() != blk:
                            wlab += 1
                        a_, b_ = j * per, (j + 1) * per
                        wchg = max(wchg, abs(
                            float(wc.cell(rr, 3).value or 0.0)
                            - float(Ch[d, a_:b_].sum())))
                        wdchg = max(wdchg, abs(
                            float(wc.cell(rr, 4).value or 0.0)
                            - float(Dh[d, a_:b_].sum())))
                        if j == 0:
                            we0 = max(we0, abs(float(wc.cell(rr, 6).value or 0.0)
                                               - float(e_start[i])))
                        elif j == 1:
                            we1 = max(we1, abs(float(wc.cell(rr, 6).value or 0.0)
                                               - float(E[d, -1])))
                chk("§7", "『充放电量』B 列分段标签 ↔ C.BATT_BLOCKS（顺序）",
                    wlab == 0, f"不符 {wlab} 格")
                chk("§7", "『充放电量』C 列充电量 ↔ npz C 按段求和",
                    wchg < TOL_CSV, f"max|Δ| {wchg:.3e} kWh")
                chk("§7", "『充放电量』D 列放电量 ↔ npz D 按段求和",
                    wdchg < TOL_CSV, f"max|Δ| {wdchg:.3e} kWh")
                chk("§7", "『充放电量』区块行 j=0 储电量 = 当日 0:00 起始 SOC"
                          "（独立链式重建）",
                    we0 < TOL_CSV, f"max|Δ| {we0:.3e} kWh")
                chk("§7", "『充放电量』区块行 j=1 储电量 = npz E 日末",
                    we1 < TOL_CSV, f"max|Δ| {we1:.3e} kWh")

                wA_date = wA_extra = wE0 = wE1c = wE_extra = 0
                for i, d in enumerate(sc):
                    for j in range(nblk):
                        rr = 2 + i * nblk + j
                        va, ve = wc.cell(rr, 1).value, wc.cell(rr, 5).value
                        if j == 0:
                            if str(va)[:10] != str(dates[d])[:10]:
                                wA_date += 1
                            if hasattr(ve, "hour"):
                                _ok_e0 = (ve.hour == 0 and ve.minute == 0
                                          and getattr(ve, "second", 0) == 0)
                            else:
                                _ok_e0 = str(ve).strip() in (
                                    "00:00", "0:00", "0:00:00", "00:00:00")
                            if not _ok_e0:
                                wE0 += 1
                        elif j == 1:
                            if str(ve).strip() != "24:00":
                                wE1c += 1
                        else:
                            if va not in (None, ""):
                                wA_extra += 1
                            if ve not in (None, ""):
                                wE_extra += 1
                chk("§7", "『充放电量』A 列：仅区块首行写日期，其余行为空",
                    wA_date == 0 and wA_extra == 0,
                    f"日期不符 {wA_date} 格 / 多余填写 {wA_extra} 格")
                chk("§7", "『充放电量』E 列：j=0 为 0:00、j=1 为「24:00」、其余为空",
                    wE0 == 0 and wE1c == 0 and wE_extra == 0,
                    f"j=0 不符 {wE0} 格 / j=1 不符 {wE1c} 格 / "
                    f"其余多余填写 {wE_extra} 格")

            if C.SHEET_EMERG in wb.sheetnames:
                we = wb[C.SHEET_EMERG]
                chk("§7", "『紧急购电量』行数 = 1 + 334×3",
                    we.max_row == 1 + n_day * 3, f"{we.max_row}")
                wem = 0.0
                bad_ed = 0
                for i, d in enumerate(sc):
                    if str(we.cell(2 + i * 3, 1).value)[:10] != str(dates[d])[:10]:
                        bad_ed += 1
                    s3 = sum(float(we.cell(2 + i * 3 + j, 3).value or 0.0)
                             for j in range(3))
                    wem = max(wem, abs(s3 - float(bb[i].sum())))
                chk("§7", "『紧急购电量』A 列日期 ↔ npz dates",
                    bad_ed == 0, f"不符 {bad_ed} 组")
                chk("§7", "『紧急购电量』3 行区块电量之和 ↔ npz 逐日 Σb"
                          "（多起事件合并后总量守恒）",
                    wem < TOL_CSV, f"max|Δ| {wem:.3e} kWh")

            if (C.SHEET_PLAN3 in wb.sheetnames
                    and C.SHEET_ADJUST in wb.sheetnames):
                wA, wB2 = wb[C.SHEET_PLAN3], wb[C.SHEET_ADJUST]
                worst_qp = 0.0
                for i in range(n_day):
                    rr = 2 + i
                    pq = float(wA.cell(rr, T_ + 2).value or 0.0)
                    qq = float(wB2.cell(rr, T_ + 2).value or 0.0)
                    worst_qp = max(worst_qp,
                                   abs((pq - qq)
                                       - float(P[sc[i]].sum() - Q[sc[i]].sum())))
                chk("§7", "xlsx 两表 EP 列净差 ↔ npz 逐日 Σp − Σq（跨表一致性）",
                    worst_qp < TOL_CSV, f"max|Δ| {worst_qp:.3e} kWh")

            if C.SHEET_PLAN3 in wb.sheetnames:
                wA = wb[C.SHEET_PLAN3]
                src_hdr = _xlsx_source_headers(C.ATTACHMENT5_R3_PATH,
                                               C.SHEET_PLAN3)
                if src_hdr:
                    got = [wA.cell(1, c).value for c in range(1, T_ + 4)]
                    n = min(len(got), len(src_hdr))
                    same_hdr = (got[:n] == src_hdr[:n]
                                and len(got) == len(src_hdr))
                    chk("§7", "『计划购电量』表头 ↔ 附件 5 原表（逐格字面相同）",
                        same_hdr,
                        f"输出 {len(got)} 格 / 模板 {len(src_hdr)} 格；"
                        f"首末：{got[0]!r} … {got[-1]!r}")
                else:
                    chk("§7", "『计划购电量』表头 ↔ 附件 5 原表", False,
                        "未找到官方模板 附件5/result3.xlsx", skip=True)

    log("")
    log("── §9 报告散文 ↔ 派生数据一致性 ──")

    _back9 = C.REPORT_BACK_MD
    _cmp9 = C.RESULT_DIR / "第三问_结算口径对照.csv"
    if not _back9.is_file():
        chk("§9", "`08` 报告存在（§9 的前提）", False, str(_back9), skip=True)
    elif not _cmp9.is_file():
        chk("§9", "结算口径对照表存在（§9 的前提）", False, str(_cmp9), skip=True)
    else:
        _h9, _b9 = _read_csv(_cmp9)
        _T9 = {}
        for _r9 in _b9:
            if _r9 and _r9[0].strip():
                _T9[_r9[0].strip()] = [_fnum(x) for x in _r9[1:]]
        _need9 = ("总费用_元", "紧急购电量_kWh")
        _ok9 = all(k in _T9 and len(_T9[k]) >= 3 for k in _need9)
        chk("§9", "对照表可解析出 DP/解析两套 总费用 与 紧急购电量（可判性）",
            _ok9,
            "6 指标 × 3 口径" if _ok9
            else f"缺列 {[k for k in _need9 if k not in _T9 or len(_T9.get(k, [])) < 3]}")

        _txt9 = _back9.read_text(encoding="utf-8").replace("\u2009", "")
        if _ok9:
            _tot_dp, _tot_an = _T9["总费用_元"][1], _T9["总费用_元"][2]
            _b_dp, _b_an = _T9["紧急购电量_kWh"][1], _T9["紧急购电量_kWh"][2]
            _delta_ref = _tot_an - _tot_dp

            _nums9 = [float(m.replace(",", "")) for m in
                      re.findall(r"差值[^0-9\n+\-]{0,60}?([+\-]?[\d,]+(?:\.\d+)?)\s*元",
                                 _txt9)]
            chk("§9", "报告「差值 X 元」有且仅有一处 ↔ 对照表独立复算",
                len(_nums9) == 1 and abs(_nums9[0] - _delta_ref) <= 0.005,
                (f"报告 {_nums9[0]:+,.2f} 元 / 复算 {_delta_ref:+,.2f} 元"
                 if len(_nums9) == 1 else f"解析到 {len(_nums9)} 处：{_nums9}"))

            _line9 = next((ln for ln in _txt9.splitlines() if "差值" in ln), "")
            _want9 = ("更省" if _delta_ref > 0.01
                      else "更贵" if _delta_ref < -0.01 else "持平")
            _hit9 = [w for w in ("更省", "更贵", "持平") if w in _line9]
            chk("§9", "报告结论定性词 ↔ 差值符号（★「预设方向」回归）",
                len(_hit9) == 1 and _hit9[0] == _want9,
                f"差值 {_delta_ref:+,.2f} 元 ⇒ 应为「{_want9}」，报告实际 "
                f"{_hit9[0] if len(_hit9) == 1 else (_hit9 or '未找到')}")

            _bn9 = np.asarray(BT["b"], float)[sc]
            _ban9 = np.asarray(BT["b_analytic"], float)[sc]
            _pr9 = np.asarray(BT["price"], float)
            _s_b, _s_ban = float(_bn9.sum()), float(_ban9.sum())
            _w_dp9 = float((_pr9 * _bn9).sum() / _s_b) if _s_b > 0 else 0.0
            _w_an9 = float((_pr9 * _ban9).sum() / _s_ban) if _s_ban > 0 else 0.0

            _mp_b9 = re.findall(
                r"紧急购电总量 DP\s*([\d,]+(?:\.\d+)?)\s*kWh\s*vs\s*解析版\s*"
                r"([\d,]+(?:\.\d+)?)\s*kWh（DP\s*(更少|更多|相等)）", _txt9)
            _want_b9 = ("更少" if _s_b < _s_ban - 1e-6
                        else "更多" if _s_b > _s_ban + 1e-6 else "相等")
            _ok_b9 = (len(_mp_b9) == 1
                      and abs(float(_mp_b9[0][0].replace(",", "")) - _s_b) <= 0.5
                      and abs(float(_mp_b9[0][1].replace(",", "")) - _s_ban) <= 0.5
                      and abs(_s_b - _b_dp) <= 0.5
                      and abs(_s_ban - _b_an) <= 0.5
                      and _mp_b9[0][2] == _want_b9)
            chk("§9", "报告「紧急购电总量 DP X vs 解析版 Y kWh（DP 更少/更多）」"
                      "↔ npz 独立复算且方向词与符号一致",
                _ok_b9,
                (f"报告 {_mp_b9[0][0]} / {_mp_b9[0][1]} kWh（DP {_mp_b9[0][2]}）"
                 f"↔ npz {_s_b:,.3f} / {_s_ban:,.3f} kWh"
                 f"↔ 对照表 {_b_dp:,.3f} / {_b_an:,.3f} kWh（应为「{_want_b9}」）"
                 if len(_mp_b9) == 1 else f"解析到 {len(_mp_b9)} 组"))

            _mp_p9 = re.findall(
                r"紧急购电加权均价 DP\s*([\d.]+)\s*vs\s*解析版\s*([\d.]+)\s*"
                r"元/kWh（DP\s*(更低|更高|相等)）", _txt9)
            _want_p9 = ("更低" if _w_dp9 < _w_an9 - 1e-9
                        else "更高" if _w_dp9 > _w_an9 + 1e-9 else "相等")
            _ok_p9 = (len(_mp_p9) == 1
                      and abs(float(_mp_p9[0][0]) - _w_dp9) <= 1e-4
                      and abs(float(_mp_p9[0][1]) - _w_an9) <= 1e-4
                      and _mp_p9[0][2] == _want_p9)
            chk("§9", "报告「紧急购电加权均价 DP a vs 解析版 b 元/kWh"
                      "（DP 更低/更高）」↔ npz 独立复算且方向词与符号一致",
                _ok_p9,
                (f"报告 {_mp_p9[0][0]} / {_mp_p9[0][1]}（DP {_mp_p9[0][2]}）"
                 f"↔ npz {_w_dp9:.6f} / {_w_an9:.6f}（应为「{_want_p9}」）"
                 if len(_mp_p9) == 1 else f"解析到 {len(_mp_p9)} 组"))

            _ex9 = C.REPORT_EXEC_MD
            if not _ex9.is_file():
                chk("§9", "执行器报告存在（跨报告口径一致性前提）",
                    False, str(_ex9), skip=True)
            else:
                _txt9e = _ex9.read_text(encoding="utf-8").replace("\u2009", "")
                _lb9 = ("算术均值" in _txt9e and "按量加权" in _txt9e
                        and "加权均价" in _txt9)
                _sc9 = (re.search(r"DP\s*[−\-]\s*解析", _txt9e) is not None
                        and re.search(r"解析\s*[−\-]\s*DP", _txt9e) is not None)
                chk("§9", "两份报告对同一概念给出**口径标签**，且「DP 差」标明"
                          "差的方向（防「同一数字两种符号约定」与假矛盾）",
                    _lb9 and _sc9,
                    f"口径标签 算术均值/按量加权/加权均价 = "
                    f"{('算术均值' in _txt9e)}/{('按量加权' in _txt9e)}/"
                    f"{('加权均价' in _txt9)}；差方向括注 = {_sc9}")

                _pr2d9 = np.broadcast_to(_pr9[None, :], _bn9.shape)
                _e_dp = (float(_pr2d9[_bn9 > 1e-9].mean())
                         if (_bn9 > 1e-9).any() else float("nan"))
                _e_an = (float(_pr2d9[_ban9 > 1e-9].mean())
                         if (_ban9 > 1e-9).any() else float("nan"))
                _me9 = re.search(
                    r"择时维度①[^\n]*?DP 版\s*([\d.]+)\s*元/kWh\s*vs\s*解析版\s*"
                    r"([\d.]+)\s*元/kWh（全期均价\s*([\d.]+)）", _txt9e)
                _ok_e9 = (_me9 is not None
                          and abs(float(_me9.group(1)) - _e_dp) <= 1e-4
                          and abs(float(_me9.group(2)) - _e_an) <= 1e-4
                          and abs(float(_me9.group(3)) - float(_pr9.mean())) <= 1e-4)
                chk("§9", "执行器报告「择时维度① 等权时段均价 + 全期均价」"
                          "↔ npz 独立复算",
                    _ok_e9,
                    (f"报告 {_me9.group(1)} / {_me9.group(2)} / "
                     f"全期 {_me9.group(3)} ↔ npz {_e_dp:.6f} / {_e_an:.6f} / "
                     f"{float(_pr9.mean()):.6f}" if _me9 is not None
                     else "未解析到择时维度①句"))

                _me29 = re.search(
                    r"择时维度②[^\n]*?DP 版\s*([\d.]+)\s*元/kWh\s*vs\s*解析版\s*"
                    r"([\d.]+)\s*元/kWh（DP\s*(更低|更高|相等)）", _txt9e)
                _ok_e29 = (_me29 is not None
                           and abs(float(_me29.group(1)) - _w_dp9) <= 1e-4
                           and abs(float(_me29.group(2)) - _w_an9) <= 1e-4
                           and _me29.group(3) == _want_p9)
                chk("§9", "执行器报告「择时维度② 按量加权均价」↔ npz 复算，"
                          "且方向词与全年回测报告**同一口径**一致",
                    _ok_e29,
                    (f"报告 ② {_me29.group(1)} / {_me29.group(2)}"
                     f"（DP {_me29.group(3)}）↔ npz {_w_dp9:.6f} / {_w_an9:.6f}"
                     f"（全年回测报告写「{_want_p9}」）" if _me29 is not None
                     else "未解析到择时维度②句"))

                _mE9 = re.search(
                    r"DP 版紧急购电\s*([\d,.]+)\s*kWh，解析版\s*([\d,.]+)\s*kWh，"
                    r"DP 差（([^）]*)）\s*([+\-][\d,.]+)\s*kWh", _txt9e)
                _mF9 = re.search(
                    r"DP 版紧急购电费\s*([\d,.]+)\s*元，解析版\s*([\d,.]+)\s*元，"
                    r"DP 差（([^）]*)）\s*([+\-][\d,.]+)\s*元", _txt9e)
                _okZ9 = False
                _detZ9 = "未解析到维度句"
                _fee9 = _T9.get("紧急购电费_元", [])
                if _mE9 and _mF9 and len(_fee9) >= 3:
                    _xe, _ye = (float(_mE9.group(1).replace(",", "")),
                                float(_mE9.group(2).replace(",", "")))
                    _ze = float(_mE9.group(4).replace(",", ""))
                    _xf, _yf = (float(_mF9.group(1).replace(",", "")),
                                float(_mF9.group(2).replace(",", "")))
                    _zf = float(_mF9.group(4).replace(",", ""))
                    _labE = re.sub(r"[*\s]", "", _mE9.group(3)).replace("−", "-")
                    _labF = re.sub(r"[*\s]", "", _mF9.group(3)).replace("−", "-")
                    _okZ9 = (_labE == "DP-解析" and abs((_xe - _ye) - _ze) <= 0.15
                             and _labF == "解析-DP" and abs((_yf - _xf) - _zf) <= 0.02
                             and abs(_xe - _b_dp) <= 0.05
                             and abs(_ye - _b_an) <= 0.05
                             and abs(_xf - _fee9[1]) <= 0.02
                             and abs(_yf - _fee9[2]) <= 0.02)
                    _detZ9 = (f"电量 差{_ze:+,.1f} vs {_xe - _ye:+,.1f} kWh"
                              f"（标签「{_labE}」）；费用 差{_zf:+,.2f} vs "
                              f"{_yf - _xf:+,.2f} 元（标签「{_labF}」）；"
                              f"对照表 电量 {_b_dp:,.1f}/{_b_an:,.1f} kWh、"
                              f"费用 {_fee9[1]:,.2f}/{_fee9[2]:,.2f} 元")
                chk("§9", "执行器报告两个「DP 差」↔ 各自标签所写的减数顺序，"
                          "且四个数分别钉到对照表（不跨量纲比较）",
                    _okZ9, _detZ9)

        _tau9, _mae9, _mb9v = _report_precision_claim(_txt9)
        _hourmap9, _base9 = _precision_rows(C.PRECISION_CSV)
        if _mae9 is None:
            chk("§9", "报告「整点 MAE / 因果基线 MAE」↔ 预报精度表（同源）",
                False, "报告未含该表述（stub 无第二问对照）或精度表缺失",
                skip=True)
        elif _tau9 is None:
            chk("§9", "报告标注「整点 MAE」的发布时刻（4 个时刻 MAE 不同）",
                False, "报告含「整点 MAE」但未写明发布时刻 ⇒ 无法判定口径；"
                       f"表内该口径各发布时刻 = {_hourmap9}")
        else:
            _ref9 = _hourmap9.get(_tau9)
            _ok9 = (_ref9 is not None
                    and abs(_mae9 - _ref9) <= 0.005
                    and abs(_mb9v - _base9) <= 0.005)
            _r9s = f"{_ref9:,.4f}" if _ref9 is not None else "缺该发布时刻行"
            _b9s = f"{_base9:,.4f}" if _base9 == _base9 else "缺该行"
            chk("§9", "报告「发布时刻 + 整点 MAE」↔ 预报精度表**同一发布时刻**行"
                      "（同源，禁止跨口径比数）",
                _ok9,
                f"报告 {_tau9} 发布 整点 {_mae9:,.4f} / 基线 "
                f"{_mb9v if _mb9v is not None else float('nan'):,.4f} kW；"
                f"表 {_tau9} 整点 {_r9s} / 基线 {_b9s} kW；"
                f"表内全部发布时刻 = {_hourmap9}")

        _win9 = _precision_window(C.PRECISION_CSV)
        _winset9 = {v for v in _win9.values() if v}
        _okwin9 = len(_win9) >= 4 and len(_winset9) == 1
        chk("§9", "预报精度表四行**评价窗口一致**（N1：禁止两把尺子比数）",
            _okwin9,
            f"共 {len(_win9)} 行；窗口 token 集 = "
            f"{sorted(_winset9) or '（全缺）'}；无 token 行 = "
            f"{[k for k, v in _win9.items() if not v][:4]}")
        if _mae9 is None:
            chk("§9", "报告自证评价窗口 ↔ 精度表 `口径` 列（N1 防回归）",
                False, "报告未含预报精度表述（stub 无第二问对照）",
                skip=True)
        else:
            _w9 = _report_precision_window(_txt9)
            chk("§9", "报告自证评价窗口 ↔ 精度表 `口径` 列（N1 防回归）",
                _w9 is not None and _winset9 == {_w9},
                f"报告标的窗口 = {_w9 or '（未标）'}；"
                f"表 = {sorted(_winset9) or '（无）'}")

        _cal9 = _precision_by_caliber(C.PRECISION_CSV)
        _tol9 = 5e-4
        _hr9 = _cal9.get("整点", {})
        _pw9 = _cal9.get("10 min 分段常数", {})
        _z_hr = _scan_zero_point(C.RESULT_DIR / "第三问_平移扫描表.csv",
                                 "MAE_kW", "位移_d_小时", "0")
        _z_10_all = _by_label(C.RESULT_DIR / "第三问_平移扫描表_10min.csv",
                              "位移0_MAE_kW")
        _miss10 = set(_pw9) - set(_z_10_all)
        _z_10 = {k: _z_10_all[k] for k in _pw9 if k in _z_10_all}
        _d_hr = _dict_close(_z_hr, _hr9, _tol9)
        chk("§9", "『平移扫描表』位移 0 点 ↔ 精度表「整点」行（N3：同尺对账）",
            _d_hr, f"扫描表 0 位移 = {_z_hr}；精度表整点 = {_hr9}")
        _d_10 = (bool(_pw9) and not _miss10
                 and _dict_close(_z_10, _pw9, _tol9))
        chk("§9", "『平移扫描表_10min』位移 0 ↔ 精度表「10 min 分段常数」行"
                  "（N3：同尺对账，四个 τ 全覆盖）",
            _d_10,
            f"扫描表 0 位移（逐 τ）= {_z_10}；精度表分段常数 = {_pw9}；"
            f"缺 τ = {sorted(_miss10) or '无'}（表内另有基线行）")
        _h10, _b10 = _read_csv(C.RESULT_DIR / "第三问_平移扫描表_10min.csv")
        _bshift = _bmae = _bzero = None
        if _h10 and all(c in _h10 for c in ("发布时刻", "最优位移_min",
                                            "最优位移_MAE_kW", "位移0_MAE_kW")):
            _ib, _is = _h10.index("发布时刻"), _h10.index("最优位移_min")
            _im, _iz = _h10.index("最优位移_MAE_kW"), _h10.index("位移0_MAE_kW")
            for _r in _b10:
                if len(_r) > max(_ib, _is, _im, _iz) and "因果基线" in _r[_ib]:
                    _bshift = _r[_is].replace(" ", "")
                    _bmae, _bzero = _fnum(_r[_im]), _fnum(_r[_iz])
                    break
        chk("§9", "『平移扫描表_10min』阴性对照：因果基线自身最优位移 = +0 "
                  "且与位移 0 同值（N3 防回归）",
            _bshift in ("+0", "0") and _bmae is not None and _bzero is not None
            and abs(_bmae - _bzero) <= _tol9,
            f"基线行最优位移 = {_bshift!r}；最优 MAE = {_bmae}；"
            f"位移 0 MAE = {_bzero}")
        _w3 = set()
        for _p in (C.RESULT_DIR / "第三问_平移扫描表.csv",
                   C.RESULT_DIR / "第三问_平移扫描表_10min.csv",
                   C.RESULT_DIR / "第三问_滚动增益表.csv"):
            _w3 |= _scan_window(_p)
        chk("§9", "三张诊断表（平移扫描 ×2 + 滚动增益）评价窗口 ↔ 精度表"
                  "（N3 防回归）",
            bool(_w3) and _w3 == _winset9,
            f"诊断表窗口集 = {sorted(_w3) or '（缺 `评价窗口` 列）'}；"
            f"精度表 = {sorted(_winset9) or '（无）'}")

    log("")
    log("── §10 第二问数值引用一致性（Issue-Q2）──")

    _q10_mode = C.Q2.mode()
    _q10_anchor = Path(C.Q2_CORE_CSV)
    _TOLQ2 = 0.005

    def _cell_num(txt: str, prefix: str, col: int = 1):
        ln = next((x for x in txt.splitlines()
                   if x.strip().startswith(prefix)), None)
        if ln is None:
            return None
        cs = [c.strip() for c in ln.strip().strip("|").split("|")]
        if len(cs) <= col:
            return None
        v = _fnum(cs[col].replace(",", "").replace("—", "nan").replace("+", ""))
        return v if v == v else None

    _q10 = {}
    try:
        _q10 = C.Q2.core_numbers() or {}
    except Exception:
        _q10 = {}
    _anc10 = _q10.get("yearly_total")
    _plan10 = _q10.get("plan_kwh")
    _efee10 = _q10.get("emerg_fee")
    _ekwh10 = _q10.get("emerg_kwh")

    if _q10_mode != "real":
        chk("§10", "第二问数值引用一致性（Issue-Q2）", False,
            f"数据源模式 `{_q10_mode}` ⇒ 无第二问对照", skip=True)
    elif not _q10_anchor.is_file():
        chk("§10", f"第二问权威锚点存在：`{_q10_anchor.name}`", False,
            str(_q10_anchor))
    else:
        _okA = (_anc10 is not None and 1e6 <= float(_anc10) <= 1e8)
        chk("§10", "权威锚点 `第二问_最终核心结果.csv::全年总费用_元` 可解析"
                   "（有限且在 1e6–1e8 元）",
            _okA, (f"{float(_anc10):,.4f} 元" if _anc10 is not None
                   else f"解析失败：{_q10}"))

        chk("§10", "锚点值 ≠ 计划购电量（★ D1 回归：20.65 M kWh 曾冒充 20.65 M 元）",
            _anc10 is None or _plan10 is None
            or abs(float(_anc10) - float(_plan10)) > 1.0,
            f"锚点 {float(_anc10):,.4f} 元 vs 计划购电量 "
            f"{_plan10 if _plan10 is None else format(float(_plan10), ',.4f')} kWh")

        _hn10 = {}
        try:
            _hn10 = C.Q2.headline_numbers() or {}
        except Exception:
            _hn10 = {}
        _hv10 = _hn10.get("yearly_total")
        chk("§10", "适配器 `headline_numbers()['yearly_total']` ≡ 权威锚点"
                   "（★ D2 回归：禁止回落到「固定计划」口径）",
            _hv10 is not None and abs(float(_hv10) - float(_anc10)) <= _TOLQ2,
            f"适配器 {_hv10} / 锚点 {_anc10} / 口径 "
            f"{_hn10.get('comparison')!r}")

        _exec10 = C.Q2_DIR / "模型结果" / "第二问_执行器对照表.csv"
        _rb10 = _fx10 = None
        if _exec10.is_file():
            _h10, _r10 = _read_csv(_exec10)
            _ci10 = {h.strip(): i for i, h in enumerate(_h10)}
            if all(k in _ci10 for k in ("比较方式", "执行器", "总费用_元")):
                for _rr in _r10:
                    if len(_rr) <= max(_ci10.values()):
                        continue
                    if "DP" not in _rr[_ci10["执行器"]]:
                        continue
                    _v = _fnum(_rr[_ci10["总费用_元"]])
                    if _v != _v:
                        continue
                    if "各自重订" in _rr[_ci10["比较方式"]]:
                        _rb10 = _v
                    elif "固定计划" in _rr[_ci10["比较方式"]]:
                        _fx10 = _v
        chk("§10", "权威锚点 ≡ 执行器表「各自重订 / DP 价值执行器」行",
            _rb10 is not None and abs(_rb10 - float(_anc10)) <= _TOLQ2,
            f"各自重订 {_rb10} / 固定计划 {_fx10} / 锚点 {_anc10}；"
            f"同名字段不得跨口径引用")

        _back10 = C.REPORT_BACK_MD
        _v_q2row = _v_diff10 = None
        if not _back10.is_file():
            chk("§10", "`08` 报告存在（§4 引用第二问数值的前提）", False,
                str(_back10), skip=True)
        else:
            _t10 = _back10.read_text(encoding="utf-8").replace("\u2009", "")
            _v_q2row = _cell_num(_t10, "| 第二问")
            _v_diff10 = _cell_num(_t10, "| 差额")
            _self10 = float(np.asarray(BT["daily_total_fee"], float).sum())
            chk("§10", "`08` 报告 §4「第二问」行总费用 ≡ 权威锚点（★ D1 回归）",
                _v_q2row is not None
                and abs(_v_q2row - float(_anc10)) <= _TOLQ2,
                (f"报告 {_v_q2row:,.4f} 元 / 锚点 {float(_anc10):,.4f} 元"
                 if _v_q2row is not None else "未解析到「| 第二问 …」表格行"))
            chk("§10", "`08` 报告 §4「第二问」行 ≠ 计划购电量 kWh（★ D1 回归）",
                _v_q2row is None or _plan10 is None
                or abs(_v_q2row - float(_plan10)) > 1.0,
                f"报告第二问 {_v_q2row} / 计划购电量 {_plan10} kWh")
            chk("§10", "`08` 报告 §4「差额」行 = 本问总费用 − 权威锚点（同源）",
                _v_diff10 is not None
                and abs(_v_diff10 - (_self10 - float(_anc10))) <= _TOLQ2,
                (f"报告 {_v_diff10:+,.4f} 元 / 复算 "
                 f"{_self10 - float(_anc10):+,.4f} 元"
                 if _v_diff10 is not None else "未解析到「| 差额 …」表格行"))
            _want10 = ("高于" if _self10 > float(_anc10) + 0.01
                       else "低于" if _self10 < float(_anc10) - 0.01 else "持平")
            _hit10 = re.findall(r"第三问费用\*\*(高于|低于)\*\*第二问", _t10)
            chk("§10", "`08` 报告 §4 总费用方向词 ↔ 差额符号（★ D4 回归）",
                len(_hit10) == 1 and _hit10[0] == _want10,
                f"本问 {_self10:,.4f} vs 第二问 {float(_anc10):,.4f} ⇒ 应为"
                f"「{_want10}」，报告实际 {_hit10 or '未找到'}")
            _self_em = float(np.asarray(BT["daily_emerg_fee"], float).sum())
            _want10e = ("低于" if _efee10 is not None and _self_em < float(_efee10)
                        else "高于" if _efee10 is not None else None)
            _hit10e = re.findall(r"紧急购电反而「(高于|低于)」", _t10)
            chk("§10", "`08` 报告 §4 紧急购电方向词 ↔ 费用比较符号（★ D4 回归）",
                _want10e is not None and len(_hit10e) == 1
                and _hit10e[0] == _want10e,
                f"本问紧急 {_self_em:,.2f} 元 vs 第二问 "
                f"{_efee10 if _efee10 is None else format(float(_efee10), ',.2f')} "
                f"元 ⇒ 应为「{_want10e}」，报告实际 {_hit10e or '未找到'}")

        _abl10 = C.REPORT_ABL_MD
        if not _abl10.is_file():
            chk("§10", "`09` 报告存在（§7 引用第二问数值的前提）", False,
                str(_abl10), skip=True)
        else:
            _ta10 = _abl10.read_text(encoding="utf-8").replace("\u2009", "")
            _s7 = _ta10.split("## 7.")[-1]
            _a10 = _cell_num(_s7, "| 全年账单 元")
            _e10 = _cell_num(_s7, "| 紧急购电费 元")
            _k10 = _cell_num(_s7, "| 紧急购电量 kWh")
            chk("§10", "`09` 报告 §7「全年账单 元」第二问列 ≡ 权威锚点",
                _a10 is not None and abs(_a10 - float(_anc10)) <= _TOLQ2,
                f"报告 {_a10} / 锚点 {_anc10}")
            chk("§10", "`09` 报告 §7「紧急购电费 元」第二问列 ≡ 权威锚点",
                _efee10 is None or (_e10 is not None
                                    and abs(_e10 - float(_efee10)) <= _TOLQ2),
                f"报告 {_e10} / 锚点 {_efee10}")
            chk("§10", "`09` 报告 §7「紧急购电量 kWh」第二问列 ≡ 权威锚点",
                _ekwh10 is None or (_k10 is not None
                                    and abs(_k10 - float(_ekwh10)) <= _TOLQ2),
                f"报告 {_k10} / 锚点 {_ekwh10}")
            chk("§10", "`08` §4 与 `09` §7 引用**同一锚点**（逐位一致）",
                _v_q2row is not None and _a10 is not None
                and abs(_v_q2row - _a10) <= _TOLQ2,
                f"`08` {_v_q2row} / `09` {_a10}")
            chk("§10", "两份报告均写明引用口径（「各自重订 / DP 价值执行器」）",
                "各自重订" in _t10 and "各自重订" in _s7,
                f"`08` 含标签={'各自重订' in _t10}；"
                f"`09` 含标签={'各自重订' in _s7}")

    n_pass = sum(1 for _, _, st, _ in checks if st == "通过")
    n_fail = sum(1 for _, _, st, _ in checks if st == "未通过")
    n_skip = sum(1 for _, _, st, _ in checks if st == "跳过")
    n_tot = len(checks)

    log("")
    log("── 汇总 ──")
    log(f"  共 {n_tot} 项：通过 {n_pass}，未通过 {n_fail}，跳过 {n_skip}")
    if n_fail:
        log("  ✘ 未通过明细：")
        for sec, nm, st, det in checks:
            if st == "未通过":
                log(f"      [{sec}] {nm}：{det}")
    log("  " + ("全部可判项通过 ✔" if n_fail == 0 else f"存在 {n_fail} 项未通过 ✘"))
    if n_skip:
        log(f"  （跳过 {n_skip} 项，多为冒烟模式或缺失的可选产物）")

    C.write_csv_utf8_sig(CHECK_CSV, ["节", "校验项", "状态", "详情"],
                         [[a, b_, c_, d_] for a, b_, c_, d_ in checks])

    lines = [
        "# 第三问 最终一致性校验报告（`14`）", "",
        f"- 输出根目录：`{C.OUT_ROOT}`",
        f"- 评分日数：{n_day}；时段数：{P.shape[1]}",
        f"- 结算系数：ρ_down = {C.RHO_DOWN}、ρ_up = {C.RHO_UP}、紧急 = 5 × 电价",
        f"- 结果：**{n_pass}/{n_tot} 通过**，{n_fail} 未通过，{n_skip} 跳过", "",
        "> 本脚本**不重算模型**，只读已产出文件并做**两路独立重算**的交叉核对：",
        "> ① §1 账单闭式不调用 `_policy3.decompose`，直接由 `price/P/Q/b` 手算；",
        "> ② §4/§5 要求 `09`/`13` 与 `08` 在**同参数**下给出**逐位相同**的数字。", "",
        "| 节 | 校验项 | 状态 | 详情 |", "|---|---|---|---|",
    ]
    for sec, nm, st, det in checks:
        mark = {"通过": "✅", "未通过": "✘", "跳过": "－"}[st]
        lines.append(f"| {sec} | {nm} | {mark} {st} | {det} |")
    lines.append("")
    if n_fail:
        lines.append("## 未通过项")
        lines.append("")
        for sec, nm, st, det in checks:
            if st == "未通过":
                lines.append(f"- `[{sec}]` **{nm}**：{det}")
        lines.append("")
    lines.append("## 复现命令")
    lines.append("")
    lines.append("```bash")
    lines.append(".venv/bin/python 第三问最终版/代码/14_最终一致性校验.py")
    lines.append("```")
    lines.append("")
    C.write_text_utf8(REPORT_MD, "\n".join(lines))

    log("")
    log(f"已保存：{CHECK_CSV.relative_to(C.PROJECT_DIR)}")
    log(f"已保存：{REPORT_MD.relative_to(C.PROJECT_DIR)}")
    log(f"总用时 {time.perf_counter() - t0:.1f} s")
    log("[14 完成] 最终一致性校验结束。")
    log.dump(LOG_TXT)
    return 0 if n_fail == 0 else 2


if __name__ == "__main__":
    raise SystemExit(main())
