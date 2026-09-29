#!/usr/bin/env python

from __future__ import annotations

import importlib.util
import shutil
import sys
import time
from datetime import datetime, time as dt_time
from pathlib import Path

import numpy as np

_HERE = Path(__file__).resolve().parent
if str(_HERE) not in sys.path:
    sys.path.insert(0, str(_HERE))


def _load(filename: str, modname: str):
    spec = importlib.util.spec_from_file_location(modname, _HERE / filename)
    mod = importlib.util.module_from_spec(spec)
    sys.modules[modname] = mod
    spec.loader.exec_module(mod)
    return mod


C = _load("_comm3.py", "q3_comm")

import openpyxl
from openpyxl.utils import get_column_letter

T = C.PERIODS_PER_DAY
ETA = C.ETA
TOLL = 1e-6
DATE_FMT = "mm-dd-yy"


def merge_emergency(b: np.ndarray, toll: float = TOLL) -> list:
    events = []
    t = 0
    while t < T:
        if b[t] > toll:
            s = t
            tot = 0.0
            while t < T and b[t] > toll:
                tot += float(b[t])
                t += 1
            events.append((s + 1, t, tot))
        else:
            t += 1
    return events


def seg_text(s: int, e: int) -> str:
    a = C.minutes_to_hhmm(C.period_bounds(s)[0])
    bb = C.minutes_to_hhmm(C.period_bounds(e)[1], use_24=True)
    return f"{a}-{bb}"


def reduce_events(events: list, k: int = 3) -> list:
    evs = [[float(s), float(e), float(tot)] for s, e, tot in events]
    while len(evs) > k:
        gaps = [evs[i + 1][0] - evs[i][1] for i in range(len(evs) - 1)]
        i = int(np.argmin(gaps))
        evs[i] = [evs[i][0], evs[i + 1][1], evs[i][2] + evs[i + 1][2]]
        del evs[i + 1]
    return evs


def main() -> int:
    C.ensure_dirs()
    log: list[str] = []

    def p(msg: str = "") -> None:
        log.append(msg)
        print(msg, flush=True)

    t_all = time.perf_counter()
    p("=" * 76)
    p("第三问 11 —— 生成并回读校验 提交结果/result3.xlsx")
    p("=" * 76)

    BT = np.load(C.BACKTEST_NPZ, allow_pickle=False)
    price = np.asarray(BT["price"], float)
    date_strs = [str(x) for x in BT["dates"]]
    score_idx = np.asarray(BT["score_day_index"], int)
    P_all = np.asarray(BT["P"], float)
    Q_all = np.asarray(BT["Q"], float)
    b_all = np.asarray(BT["b"], float)
    C_all = np.asarray(BT["C"], float)
    D_all = np.asarray(BT["D"], float)
    E_all = np.asarray(BT["E"], float)
    yearly_total = float(np.asarray(BT["yearly_total"], float).ravel()[0])

    E_start = np.zeros(365)
    _prev = float(C.E_INIT)
    for d in score_idx:
        E_start[int(d)] = _prev
        _prev = float(E_all[int(d), -1])

    days = [int(d) for d in score_idx]
    day_dates = [date_strs[d] for d in days]

    U_all = np.maximum(P_all - Q_all, 0.0)
    V_all = np.maximum(Q_all - P_all, 0.0)
    plan_fee = (price * np.minimum(P_all, Q_all)).sum(axis=1)
    down_fee = (C.RHO_DOWN * price * U_all).sum(axis=1)
    up_fee = (C.RHO_UP * price * V_all).sum(axis=1)
    emerg_fee = (C.EMERG_MULT * price * b_all).sum(axis=1)

    p(f"数据源：{C.BACKTEST_NPZ.name}（两阶段 SP 计划 + 滚动调整 + 阶段内 DP 执行）")
    p(f"提交 {len(days)} 天：{day_dates[0]} … {day_dates[-1]}")
    p(f"全年账单（08 口径）= {yearly_total:,.6f} 元；"
      f"计划 {plan_fee[score_idx].sum():,.2f} + 下调 "
      f"{down_fee[score_idx].sum():,.2f} + 上调 {up_fee[score_idx].sum():,.2f} + 紧急 "
      f"{emerg_fee[score_idx].sum():,.2f}")

    C.SUBMIT_DIR.mkdir(parents=True, exist_ok=True)
    out_path = C.RESULT3_XLSX
    shutil.copyfile(C.ATTACHMENT5_R3_PATH, out_path)
    p(f"已复制模板：{C.ATTACHMENT5_R3_PATH.name} → {out_path}")

    wb = openpyxl.load_workbook(out_path)
    p(f"工作表：{wb.sheetnames}")

    ws1 = wb[C.SHEET_PLAN3]
    hdr1 = [ws1.cell(1, c).value for c in range(1, ws1.max_column + 1)]
    ncol = len(hdr1)
    p("")
    p(f"── 工作表『{C.SHEET_PLAN3}』：{ws1.max_row} 行 × {ncol} 列（表头保持原样）──")
    p(f"  表头首列 = {hdr1[0]!r}；第 2 列 = {hdr1[1]!r}；"
      f"第 {ncol - 2} 列 = {hdr1[ncol - 3]!r}；末两列 = {hdr1[-2]!r}, {hdr1[-1]!r}")
    assert ncol == T + 3, f"意外的列数：{ncol}"
    ep_col, eq_col = T + 2, T + 3

    for i, d in enumerate(days):
        row = 2 + i
        av = ws1.cell(row, 1).value
        av_s = av.strftime("%Y-%m-%d") if hasattr(av, "strftime") else str(av)
        if av_s[:10] != day_dates[i]:
            p(f"  ⚠ 第 {row} 行日期不符：模板 {av_s} vs 预期 {day_dates[i]}，已覆盖写入。")
            ws1.cell(row, 1).value = datetime.strptime(day_dates[i], "%Y-%m-%d")
            ws1.cell(row, 1).number_format = DATE_FMT
        for t in range(T):
            ws1.cell(row, 2 + t).value = float(P_all[d, t])
        ws1.cell(row, ep_col).value = float(P_all[d].sum())
        ws1.cell(row, eq_col).value = float(plan_fee[d])
    p(f"  已填写第 2..{1 + len(days)} 行的列 B..{get_column_letter(T + 1)}"
      f"（{T} 个时段）与 {get_column_letter(ep_col)}(全天计划购电量)、"
      f"{get_column_letter(eq_col)}(计划购电费 Σc·min(p,q))")

    ws2 = wb[C.SHEET_ADJUST]
    hdr2 = [ws2.cell(1, c).value for c in range(1, ws2.max_column + 1)]
    p("")
    p(f"── 工作表『{C.SHEET_ADJUST}』：{ws2.max_row} 行 × {len(hdr2)} 列（表头保持原样）──")
    assert len(hdr2) == T + 3, f"意外的列数：{len(hdr2)}"
    for i, d in enumerate(days):
        row = 2 + i
        av = ws2.cell(row, 1).value
        av_s = av.strftime("%Y-%m-%d") if hasattr(av, "strftime") else str(av)
        if av_s[:10] != day_dates[i]:
            p(f"  ⚠ 第 {row} 行日期不符：模板 {av_s} vs 预期 {day_dates[i]}，已覆盖写入。")
            ws2.cell(row, 1).value = datetime.strptime(day_dates[i], "%Y-%m-%d")
            ws2.cell(row, 1).number_format = DATE_FMT
        for t in range(T):
            ws2.cell(row, 2 + t).value = float(Q_all[d, t])
        ws2.cell(row, ep_col).value = float(Q_all[d].sum())
        ws2.cell(row, eq_col).value = float(down_fee[d] + up_fee[d])
    p(f"  已填写第 2..{1 + len(days)} 行的列 B..{get_column_letter(T + 1)}"
      f"（{T} 个时段的调整购电量 q = p − u + v）与 "
      f"{get_column_letter(ep_col)}(全天调整购电量)、"
      f"{get_column_letter(eq_col)}(调整相关费用 0.5Σc·u + 1.5Σc·v)")

    ws3 = wb[C.SHEET_BATT]
    hdr3 = [ws3.cell(1, c).value for c in range(1, 7)]
    p("")
    p(f"── 工作表『{C.SHEET_BATT}』：重建 334 个 6 行区块（表头 {hdr3}）──")
    if ws3.max_row >= 2:
        ws3.delete_rows(2, ws3.max_row - 1)
    r = 2
    for i, d in enumerate(days):
        base = r + 6 * i
        dtv = datetime.strptime(day_dates[i], "%Y-%m-%d")
        for j in range(6):
            rr = base + j
            ws3.cell(rr, 2).value = C.BATT_BLOCKS[j]
            a, bb = j * 24, (j + 1) * 24
            ws3.cell(rr, 3).value = float(C_all[d, a:bb].sum())
            ws3.cell(rr, 4).value = float(D_all[d, a:bb].sum())
            if j == 0:
                c = ws3.cell(rr, 1)
                c.value = dtv
                c.number_format = DATE_FMT
                c = ws3.cell(rr, 5)
                c.value = dt_time(0, 0)
                c.number_format = "h:mm"
                ws3.cell(rr, 6).value = float(E_start[d])
            elif j == 1:
                ws3.cell(rr, 5).value = "24:00"
                ws3.cell(rr, 6).value = float(E_all[d, -1])
    p(f"  已写入第 2..{1 + 6 * len(days)} 行（{len(days)} 天 × 6 段）")
    p("  列 E/F 仅在每个区块前两行填 时刻 与 储电量（0:00 / 24:00）")

    ws4 = wb[C.SHEET_EMERG]
    hdr4 = [ws4.cell(1, c).value for c in range(1, 4)]
    p("")
    p(f"── 工作表『{C.SHEET_EMERG}』：重建 334 个 3 行区块（表头 {hdr4}）──")
    if ws4.max_row >= 2:
        ws4.delete_rows(2, ws4.max_row - 1)
    r = 2
    n_ev_raw = 0
    n_merged = 0
    for i, d in enumerate(days):
        raw = merge_emergency(b_all[d])
        n_ev_raw += len(raw)
        ev = reduce_events(raw, 3)
        if len(raw) > 3:
            n_merged += 1
        dtv = datetime.strptime(day_dates[i], "%Y-%m-%d")
        for j in range(3):
            rr = r + j
            if j == 0:
                c = ws4.cell(rr, 1)
                c.value = dtv
                c.number_format = DATE_FMT
            if j < len(ev):
                s, e, tot = ev[j]
                ws4.cell(rr, 2).value = seg_text(int(round(s)), int(round(e)))
                ws4.cell(rr, 3).value = float(tot)
            else:
                ws4.cell(rr, 2).value = ""
        r += 3
    p(f"  已写入第 2..{r - 1} 行（{len(days)} 天 × 3 行）")
    p(f"  原始事件合计 {n_ev_raw} 起；因原始事件 > 3 起而需合并呈现的天数：{n_merged}")

    wb.save(out_path)
    p(f"已保存：{out_path}")

    p("")
    p("── 回读校验 ──")
    wb2 = openpyxl.load_workbook(out_path, data_only=True)
    checks = []

    def chk(name, ok, detail=""):
        checks.append((name, bool(ok), detail))
        p(f"  {'✔' if ok else '✘'} {name}{('：' + detail) if detail else ''}")

    wsA = wb2[C.SHEET_PLAN3]
    chk("计划购电量 行数 = 335", wsA.max_row == 335, f"实际 {wsA.max_row}")
    chk("计划购电量 列数 = 147", wsA.max_column == 147, f"实际 {wsA.max_column}")
    chk("计划购电量 表头首列/末两列不变",
        wsA.cell(1, 1).value == hdr1[0] and wsA.cell(1, T + 2).value == hdr1[-2]
        and wsA.cell(1, T + 3).value == hdr1[-1])
    ok_date = True
    for i, dtxt in enumerate(day_dates):
        v = wsA.cell(2 + i, 1).value
        if not hasattr(v, "strftime") or v.strftime("%Y-%m-%d") != dtxt:
            ok_date = False
            break
    chk("计划购电量 列 A 日期 = 2025-02-01..2025-12-31 连续", ok_date)

    worst_p, worst_pq, worst_pc = 0.0, 0.0, 0.0
    for i, d in enumerate(days):
        row = 2 + i
        for t in range(T):
            worst_p = max(worst_p, abs(float(wsA.cell(row, 2 + t).value) - P_all[d, t]))
        worst_pq = max(worst_pq, abs(float(wsA.cell(row, T + 2).value) - P_all[d].sum()))
        worst_pc = max(worst_pc, abs(float(wsA.cell(row, T + 3).value) - plan_fee[d]))
    chk("计划购电量 144 时段逐格与模型一致", worst_p < 1e-6, f"最大偏差 {worst_p:.3e}")
    chk("计划购电量 全天购电量列自洽", worst_pq < 1e-6, f"最大偏差 {worst_pq:.3e}")
    chk("计划购电量 全天购电费列自洽", worst_pc < 1e-6, f"最大偏差 {worst_pc:.3e}")

    wsB = wb2[C.SHEET_ADJUST]
    chk("调整购电量 行数 = 335", wsB.max_row == 335, f"实际 {wsB.max_row}")
    chk("调整购电量 列数 = 147", wsB.max_column == 147, f"实际 {wsB.max_column}")
    worst_q, worst_qq, worst_qc = 0.0, 0.0, 0.0
    for i, d in enumerate(days):
        row = 2 + i
        for t in range(T):
            worst_q = max(worst_q, abs(float(wsB.cell(row, 2 + t).value) - Q_all[d, t]))
        worst_qq = max(worst_qq, abs(float(wsB.cell(row, T + 2).value) - Q_all[d].sum()))
        worst_qc = max(worst_qc, abs(float(wsB.cell(row, T + 3).value)
                                     - (down_fee[d] + up_fee[d])))
    chk("调整购电量 144 时段逐格与模型一致", worst_q < 1e-6, f"最大偏差 {worst_q:.3e}")
    chk("调整购电量 全天购电量列自洽", worst_qq < 1e-6, f"最大偏差 {worst_qq:.3e}")
    chk("调整购电量 全天费用列自洽", worst_qc < 1e-6, f"最大偏差 {worst_qc:.3e}")

    worst_neq = 0.0
    sum_abs_pq = 0.0
    for i, d in enumerate(days):
        row = 2 + i
        for t in range(T):
            pv_ = float(wsA.cell(row, 2 + t).value)
            qv_ = float(wsB.cell(row, 2 + t).value)
            sum_abs_pq += abs(pv_ - qv_)
            if pv_ - qv_ > TOLL and U_all[d, t] <= TOLL:
                worst_neq = max(worst_neq, abs(pv_ - qv_))
            if qv_ - pv_ > TOLL and V_all[d, t] <= TOLL:
                worst_neq = max(worst_neq, abs(qv_ - pv_))
    chk("计划/调整两表的差额与 u/v 分解符号一致", worst_neq < TOLL,
        f"冲突量最大 {worst_neq:.3e} kWh；Σ|p−q| = {sum_abs_pq:,.4f} kWh"
        f"（= Σu + Σv = {U_all[score_idx].sum() + V_all[score_idx].sum():,.4f}）")

    wsC = wb2[C.SHEET_BATT]
    chk("充放电量 行数 = 1 + 334×6", wsC.max_row == 1 + 6 * len(days),
        f"实际 {wsC.max_row}")
    ok_blk = True
    worst_batt = 0.0
    for i, d in enumerate(days):
        base = 2 + 6 * i
        for j in range(6):
            rr = base + j
            if wsC.cell(rr, 2).value != C.BATT_BLOCKS[j]:
                ok_blk = False
            a, bb = j * 24, (j + 1) * 24
            worst_batt = max(worst_batt,
                             abs(float(wsC.cell(rr, 3).value) - C_all[d, a:bb].sum()),
                             abs(float(wsC.cell(rr, 4).value) - D_all[d, a:bb].sum()))
    chk("充放电量 六段标签与行序正确", ok_blk)
    chk("充放电量 充电/放电量逐格与模型一致", worst_batt < 1e-6,
        f"最大偏差 {worst_batt:.3e}")
    ok_ef = True
    worst_soc = 0.0
    for i, d in enumerate(days):
        base = 2 + 6 * i
        e0 = wsC.cell(base, 6).value
        e1 = wsC.cell(base + 1, 6).value
        e0 = float(e0) if e0 is not None else np.nan
        e1 = float(e1) if e1 is not None else np.nan
        worst_soc = max(worst_soc, abs(e0 - E_start[d]), abs(e1 - E_all[d, -1]))
        if wsC.cell(base + 1, 5).value != "24:00":
            ok_ef = False
    chk("充放电量 列 E 的 0:00 / 24:00 标签与列 F 储电量自洽", ok_ef,
        f"储电量最大偏差 {worst_soc:.3e}")

    worst_rec = 0.0
    for i, d in enumerate(days):
        worst_rec = max(worst_rec,
                        abs((ETA * C_all[d].sum() - D_all[d].sum() / ETA)
                            - (E_all[d, -1] - E_start[d])))
    chk("储能 SOC 递推自洽（ηΣC − ΣD/η = ΔE）", worst_rec < 1e-6,
        f"最大残差 {worst_rec:.3e}")

    chk("储能 SOC 全程落在 [1200, 10800] kWh",
        float(E_all[score_idx].min()) >= C.E_MIN - 1e-6
        and float(E_all[score_idx].max()) <= C.E_MAX + 1e-6,
        f"范围 [{E_all[score_idx].min():.3f}, {E_all[score_idx].max():.3f}]")

    chk("单时段充放电量 ≤ S = 5000·Δt",
        float(C_all[score_idx].max()) <= C.S_PERIOD_KWH + 1e-6
        and float(D_all[score_idx].max()) <= C.S_PERIOD_KWH + 1e-6,
        f"最大充 {C_all[score_idx].max():.6f}，最大放 {D_all[score_idx].max():.6f} kWh")

    wsD = wb2[C.SHEET_EMERG]
    chk("紧急购电量 行数 = 1 + 334×3", wsD.max_row == 1 + 3 * len(days),
        f"实际 {wsD.max_row}")
    ok_dt4 = True
    for i, dtxt in enumerate(day_dates):
        v = wsD.cell(2 + 3 * i, 1).value
        if not hasattr(v, "strftime") or v.strftime("%Y-%m-%d") != dtxt:
            ok_dt4 = False
            break
    chk("紧急购电量 每个区块首行日期正确", ok_dt4)
    tot_ev = sum(float(b_all[d].sum()) for d in days)
    tot_in = 0.0
    for rr in range(2, wsD.max_row + 1):
        v = wsD.cell(rr, 3).value
        if isinstance(v, (int, float)):
            tot_in += float(v)
    chk("紧急购电量 求和 = 模型紧急量合计", abs(tot_in - tot_ev) < 1e-4,
        f"表内 {tot_in:.6f} vs 模型 {tot_ev:.6f} kWh")

    fee_tbl = (sum(float(wsA.cell(2 + i, T + 3).value) for i in range(len(days)))
               + sum(float(wsB.cell(2 + i, T + 3).value) for i in range(len(days)))
               + emerg_fee[score_idx].sum())
    chk("四表费用闭合（计划表 + 调整表 + 紧急费 = 全年账单）",
        abs(fee_tbl - yearly_total) < 1e-3,
        f"表内 {fee_tbl:.6f} vs 08 结算 {yearly_total:.6f} 元")

    Z = C.Q2.matrix()
    load_e = np.asarray(Z["load_energy_kwh"], float)[score_idx]
    pv_e = np.asarray(Z["pv_energy_kwh"], float)[score_idx]
    U_all2 = np.asarray(BT["U"], float)[score_idx]
    lhs = Q_all[score_idx].sum() + b_all[score_idx].sum() + D_all[score_idx].sum()
    rhs = load_e.sum() - pv_e.sum() + C_all[score_idx].sum() + U_all2.sum()
    chk("全年电量守恒（q + b + D = 负荷 − 光伏 + 充电 + 弃电）",
        abs(lhs - rhs) < 1e-2, f"LHS {lhs:,.4f} vs RHS {rhs:,.4f} kWh")

    yr_path = C.YEARLY_CSV
    if yr_path.exists():
        import csv
        with open(yr_path, encoding="utf-8-sig") as fh:
            yr = {r["指标"]: r["数值"] for r in csv.DictReader(fh)}
        ref = None
        for k in yr:
            if "总费用" in k and "元" in k:
                ref = float(yr[k])
                break
        if ref is not None:
            chk("与 08 年度汇总表总费用一致", abs(ref - yearly_total) < 1e-3,
                f"年度表 {ref:.6f} vs NPZ {yearly_total:.6f} 元")
    else:
        p(f"  （未找到 {yr_path.name}，跳过与年度汇总表的一致性核对）")

    n_fail = sum(1 for _, ok, _ in checks if not ok)
    p("")
    p(f"校验结果：{len(checks) - n_fail}/{len(checks)} 通过"
      f"{'，全部通过 ✔' if n_fail == 0 else f'，失败 {n_fail} 项 ✘'}")
    p(f"提交文件：{out_path}（{out_path.stat().st_size / 1024:.1f} KiB）")
    p(f"总用时 {time.perf_counter() - t_all:.1f} s")

    rep = ["# 第三问 result3.xlsx 生成与回读校验报告", "",
           f"- 代码：`第三问最终版/代码/11_生成并校验result3.py`",
           f"- 提交文件：`第三问最终版/提交结果/result3.xlsx`"
           f"（{out_path.stat().st_size / 1024:.1f} KiB）",
           f"- 评分期 {len(days)} 天：{day_dates[0]} … {day_dates[-1]}",
           f"- 全年账单（08 结算口径）= **{yearly_total:,.6f} 元**", "",
           "## 1. 四张工作表的口径", "",
           "| 工作表 | 尺寸 | 内容 |",
           "|---|---|---|",
           f"| `{C.SHEET_PLAN3}` | 335 × 147 | 144 个时段的计划购电量 $p_t$；"
           "列 EP = $\\sum p_t$（全天计划购电量）；列 EQ = $\\sum c_t\\min(p_t,q_t)$ |",
           f"| `{C.SHEET_ADJUST}` | 335 × 147 | 144 个时段的调整购电量 "
           "$q_t=p_t-u_t+v_t$；列 EP = $\\sum q_t$；列 EQ = "
           "$0.5\\sum c_t u_t + 1.5\\sum c_t v_t$ |",
           f"| `{C.SHEET_BATT}` | 2005 × 6 | 334 天 × 6 个 4 h 分段；"
           "列 C = 充电量 $C_t$、列 D = 放电量 $D_t$（交流侧）；"
           "列 E/F 仅在区块前两行填 `0:00`/`24:00` 与储电量 |",
           f"| `{C.SHEET_EMERG}` | 1003 × 3 | 334 天 × 3 行；"
           "连续 `b_t>0` 的时段合并为事件 |",
           "",
           "**费用闭合**（严格照 `C题.pdf` 第 2 页费率条款）：", "",
           "$$K_d=\\sum_t c_t\\min(p_t,q_t)"
           "+\\sum_t\\big[0.5c_t(p_t-q_t)^+ + 1.5c_t(q_t-p_t)^+\\big]"
           "+\\sum_t 5c_t b_t$$", "",
           f"四表相加 = {fee_tbl:,.6f} 元；08 结算 = {yearly_total:,.6f} 元；"
           f"差 {fee_tbl - yearly_total:+.3e} 元。", "",
           "## 2. 回读校验结果", "",
           "| # | 校验项 | 结果 | 说明 |",
           "|---:|---|---|---|"]
    for k, (nm, ok, det) in enumerate(checks, 1):
        rep.append(f"| {k} | {nm} | {'✔' if ok else '✘'} | {det} |")
    rep += ["", f"**{len(checks) - n_fail}/{len(checks)} 通过"
            f"{'，全部通过 ✔' if n_fail == 0 else f'，失败 {n_fail} 项 ✘'}**", "",
            "## 3. 口径提示", "",
            "* 时间标签沿用官方模板的**整体错位**：第 $N$ 个数据列对应模型第 $N$ 个"
            "时段 $[(N-1)\\Delta t, N\\Delta t)$；表头文字**未被改写**。",
            "* 紧急购电量只可能来自**实际执行结果**，且**严禁用于主动充电**"
            "（README A-30）；本问全年实测「紧急购电与充电同段」天数 = 0。",
            "* `充放电量` 表为**交流侧**口径，故 SOC 递推按 "
            "$\\eta\\Sigma C-\\Sigma D/\\eta=\\Delta E$ 校验。",
            ""]
    C.write_text_utf8(C.REPORT_RESULT3_MD, "\n".join(rep))
    p("已保存：报告/第三问_result3生成与校验报告.md")
    C.write_text_utf8(C.SOLVE_LOG_DIR / "第三问_11结果生成日志.txt",
                      "\n".join(log + ["", "[11 完成] result3.xlsx 生成与回读校验结束。"]))
    p("已保存：求解日志/第三问_11结果生成日志.txt")
    return 0 if n_fail == 0 else 1


if __name__ == "__main__":
    raise SystemExit(main())
